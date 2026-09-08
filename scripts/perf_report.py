"""How long the product takes, on the decks it has and on decks bigger than any.

    python scripts/perf_report.py             # every deck on this machine
    python scripts/perf_report.py --large     # and synthesised 100/200/400-slide decks
    python scripts/perf_report.py --json

Measured rather than estimated, and it has already paid for itself twice. The
first run found the chart and SmartArt guards rescanning the whole package once
per change -- a 1,199-change tidy took **7 minutes 48 seconds** -- and the second
found `Package.read` reopening the archive on every part read, which was 84% of
`inspect`. Neither was a bug anyone had reported. Nothing was ever *wrong*; it
was correct, verified, and took eight minutes.

The largest real deck available is 52 slides, so `--large` synthesises the sizes
nobody has a fixture for. Those decks are text only and therefore light per
slide: they measure how the engine scales with *shape count*, and the corpus
decks measure what media and charts cost. Both numbers are needed and neither
substitutes for the other.
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src" / "engine"))

from slide_wright.audit import audit as audit_deck  # noqa: E402
from slide_wright.brand import plan_conformance, read_profile  # noqa: E402
from slide_wright.diff import diff  # noqa: E402
from slide_wright.inspect import inspect  # noqa: E402
from slide_wright.layout import plan_alignment  # noqa: E402
from slide_wright.session import Session  # noqa: E402

CORPUS = ROOT / "tests" / "corpus"
FIXTURES = ROOT / "tests" / "fixtures" / "third-party"


def took(fn):
    start = time.perf_counter()
    value = fn()
    return value, time.perf_counter() - start


def measure(deck: Path, workspace: Path) -> dict:
    """Open, read, audit, plan a tidy, apply it, and diff the result."""
    session, t_open = took(lambda: Session.open(deck, workspace=workspace))
    info, t_inspect = took(session.deck)
    _, t_audit = took(lambda: audit_deck(info, deck.name))

    def plan():
        changeset = session.propose("tidy")
        try:
            for change in plan_conformance(info, read_profile(deck)).changes:
                changeset.add(change)
        except Exception:
            pass
        for change in plan_alignment(info).changes:
            changeset.add(change)
        return changeset

    changeset, t_plan = took(plan)
    changeset.approve_all()
    changes = len(changeset.approved)

    t_apply = t_diff = None
    if changes:
        try:
            _, t_apply = took(lambda: session.apply("tidy"))
            _, t_diff = took(
                lambda: diff(session.versions[0].path, session.current.path)
            )
        except Exception:
            t_apply = t_diff = None

    return {
        "deck": deck.name,
        "mb": round(deck.stat().st_size / 1024 / 1024, 1),
        "slides": len(info.slides),
        "shapes": sum(len(s.shapes) for s in info.slides),
        "changes": changes,
        "open_s": round(t_open, 3),
        "inspect_s": round(t_inspect, 3),
        "audit_s": round(t_audit, 3),
        "plan_s": round(t_plan, 3),
        "apply_s": round(t_apply, 3) if t_apply is not None else None,
        "diff_s": round(t_diff, 3) if t_diff is not None else None,
    }


def synthesise(slides: int, per_slide: int, path: Path) -> Path:
    """A deck larger than any real one here. Text only, so it measures shapes."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for n in range(slides):
        slide = prs.slides.add_slide(blank)
        for i in range(per_slide):
            box = slide.shapes.add_textbox(
                Inches(0.4 + (i % 4) * 3.0), Inches(0.5 + (i // 4) * 0.75),
                Inches(2.8), Inches(0.6),
            )
            run = box.text_frame.paragraphs[0].add_run()
            run.text = f"Slide {n + 1} item {i + 1}: revenue 12.4x, margin 21.{i}%"
            run.font.name = "Arial"
            run.font.size = Pt(12)
    prs.save(path)
    return path


HEADER = (
    f"  {'deck':<32}{'slides':>7}{'shapes':>8}{'MB':>6}"
    f"{'open':>7}{'inspect':>9}{'audit':>8}{'plan':>7}{'apply':>8}{'diff':>7}"
)


def line(row: dict) -> str:
    def fmt(key):
        value = row[key]
        return f"{value:.2f}s" if value is not None else "     —"

    return (
        f"  {row['deck'][:31]:<32}{row['slides']:>7}{row['shapes']:>8}{row['mb']:>6}"
        f"{fmt('open_s'):>7}{fmt('inspect_s'):>9}{fmt('audit_s'):>8}"
        f"{fmt('plan_s'):>7}{fmt('apply_s'):>8}{fmt('diff_s'):>7}"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--large", action="store_true",
                        help="also synthesise 100, 200 and 400-slide decks")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    tmp = Path(tempfile.mkdtemp())
    decks = sorted(CORPUS.glob("*.pptx"))
    if FIXTURES.is_dir():
        decks += sorted(FIXTURES.glob("*.pptx"))

    rows = []
    for deck in decks:
        staged = tmp / deck.name
        staged.write_bytes(deck.read_bytes())
        try:
            rows.append(measure(staged, tmp / f"ws-{deck.stem}"))
        except Exception as exc:  # a deck that will not open is not a timing
            print(f"skipped {deck.name}: {exc}", file=sys.stderr)

    synthetic = []
    if args.large:
        for slides, per_slide in ((100, 8), (200, 8), (400, 6)):
            path = synthesise(slides, per_slide, tmp / f"synthetic-{slides}-slides.pptx")
            synthetic.append(measure(path, tmp / f"ws-big-{slides}"))

    if args.json:
        print(json.dumps({"corpus": rows, "synthetic": synthetic}, indent=2))
        return 0

    print("PIPELINE TIMINGS\n")
    print(HEADER)
    for row in sorted(rows, key=lambda r: -r["shapes"]):
        print(line(row))

    if synthetic:
        print("\n  larger than any real deck here (text only, so this is shape scaling)\n")
        for row in synthetic:
            print(line(row))
        applied = [(r["changes"], r["apply_s"]) for r in synthetic if r["apply_s"]]
        if len(applied) > 1:
            per = [t / n for n, t in applied]
            print(
                f"\n  apply cost per change: "
                f"{min(per) * 1000:.2f}–{max(per) * 1000:.2f} ms "
                f"(median {statistics.median(per) * 1000:.2f}) across "
                f"{applied[0][0]}–{applied[-1][0]} changes"
            )
            print("  Linear in the number of changes is the property worth having.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
