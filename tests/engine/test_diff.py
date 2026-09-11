"""Structural diff — what changed, in a reviewer's terms.

The case that motivates this module: an edit produces a change nobody asked
for, the deck is blocked, and the reviewer is handed a part name. Everything
here exists so that moment produces a sentence instead of a filename.
"""

from __future__ import annotations

import shutil

import pytest

from slide_wright.apply import apply_changes
from slide_wright.changeset import Change, ChangeSet, Op
from slide_wright.diff import VISIBLE_MOVE_EMU, ShapeDelta, _describe_text_change, diff
from slide_wright.inspect import inspect


def approved(deck, *changes: Change) -> ChangeSet:
    cs = ChangeSet(deck=str(deck))
    for c in changes:
        cs.add(c)
    cs.approve_all()
    return cs


@pytest.fixture
def table(adversarial_deck):
    return next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")


@pytest.fixture
def positioned(adversarial_deck):
    """A plain shape the applier will actually move.

    Not a placeholder -- those are positioned by the layout and moving one is
    refused. Not a chart -- read-only under ADR-0009. Either would make this
    fixture yield a shape whose move silently does not happen, and the diff
    assertions would then be testing the refusal rather than the diff.
    """
    return next(
        s for s in inspect(adversarial_deck).slides[1].shapes
        if s.kind == "shape" and s.x is not None and not s.geometry_inherited
    )


class TestNoDifference:
    def test_a_copy_has_no_structural_differences(self, adversarial_deck, tmp_path):
        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        result = diff(adversarial_deck, copy)
        assert not result.changed
        assert "No structural differences" in result.render()

    def test_a_deck_does_not_differ_from_itself(self, adversarial_deck):
        assert not diff(adversarial_deck, adversarial_deck).changed


class TestTextChanges:
    def test_a_cell_edit_is_reported_as_a_text_change(
        self, adversarial_deck, table, tmp_path
    ):
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table.id}/r1/c1", before="9.4x", after="11.8x")), out)

        deltas = diff(adversarial_deck, out).deltas
        assert len(deltas) == 1
        assert deltas[0].kind == "text"
        assert deltas[0].slide == 3
        assert deltas[0].is_content

    def test_the_description_shows_the_changed_region_not_the_string_head(
        self, adversarial_deck, table, tmp_path
    ):
        """A table's text is long; truncating both sides shows the same prefix twice."""
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table.id}/r1/c1", before="9.4x", after="11.8x")), out)

        description = diff(adversarial_deck, out).deltas[0].description
        assert "9.4" in description and "11.8" in description, description


class TestDescribeTextChange:
    """The helper directly, because its whole value is in the wording."""

    def test_isolates_a_change_inside_a_long_string(self):
        before = "CompanyEV/EBITDAMarginGrowthAlpha Corp9.4x22.1%18%Beta Industries"
        after = "CompanyEV/EBITDAMarginGrowthAlpha Corp11.8x22.1%18%Beta Industries"
        out = _describe_text_change(before, after)
        assert "9.4 -> 11.8" in out
        assert len(out) < len(before), "the point is to be shorter than the string"

    def test_a_short_change_needs_no_context(self):
        assert _describe_text_change("Q3 results", "Q4 results") == "Q[3 -> 4] results"

    def test_a_wholly_different_string_is_shown_whole(self):
        assert _describe_text_change("short", "other") == "[short -> other]"

    def test_deletion_and_insertion_are_named(self):
        assert "(nothing)" in _describe_text_change("text", "")
        assert "(nothing)" in _describe_text_change("", "text")

    def test_newlines_do_not_break_the_line(self):
        assert "\n" not in _describe_text_change("a\nb", "a\nc")


class TestGeometryChanges:
    def test_a_move_is_reported_in_inches(self, adversarial_deck, positioned, tmp_path):
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.MOVE, slide=2, target=positioned.id,
            before=(positioned.x, positioned.y),
            after=(positioned.x + 457200, positioned.y))), out)

        deltas = diff(adversarial_deck, out).deltas
        assert [d.kind for d in deltas] == ["geometry"]
        assert "+0.50in" in deltas[0].description

    def test_a_move_is_not_a_content_change(self, adversarial_deck, positioned, tmp_path):
        """Moving a box changes how the deck looks, not what it says."""
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.MOVE, slide=2, target=positioned.id,
            before=(positioned.x, positioned.y),
            after=(positioned.x + 457200, positioned.y))), out)

        result = diff(adversarial_deck, out)
        assert result.deltas and not result.content_deltas

    def test_a_sub_threshold_move_is_flagged_as_probable_rounding(
        self, adversarial_deck, positioned, tmp_path
    ):
        """Reported, not hidden — but named so nobody chases a save artefact."""
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.MOVE, slide=2, target=positioned.id,
            before=(positioned.x, positioned.y),
            after=(positioned.x + VISIBLE_MOVE_EMU - 1, positioned.y))), out)

        deltas = diff(adversarial_deck, out).deltas
        assert deltas, "a sub-threshold move must still be reported"
        assert "rounding" in deltas[0].description

    def test_a_resize_is_reported_with_both_sizes(
        self, adversarial_deck, positioned, tmp_path
    ):
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c1", op=Op.RESIZE, slide=2, target=positioned.id,
            before=(positioned.cx, positioned.cy),
            after=(positioned.cx // 2, positioned.cy))), out)

        deltas = [d for d in diff(adversarial_deck, out).deltas if d.kind == "size"]
        assert deltas and "->" in deltas[0].description


class TestRendering:
    def test_groups_differences_by_slide(self, adversarial_deck, table, positioned, tmp_path):
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(
            adversarial_deck,
            Change(id="c1", op=Op.SET_TABLE_CELL, slide=3,
                   target=f"{table.id}/r1/c1", before="9.4x", after="11.8x"),
            Change(id="c2", op=Op.MOVE, slide=2, target=positioned.id,
                   before=(positioned.x, positioned.y),
                   after=(positioned.x + 457200, positioned.y)),
        ), out)

        out_text = diff(adversarial_deck, out).render()
        assert "slide 2" in out_text and "slide 3" in out_text

    def test_separates_what_it_says_from_how_it_looks(
        self, adversarial_deck, table, positioned, tmp_path
    ):
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, approved(
            adversarial_deck,
            Change(id="c1", op=Op.SET_TABLE_CELL, slide=3,
                   target=f"{table.id}/r1/c1", before="9.4x", after="11.8x"),
            Change(id="c2", op=Op.MOVE, slide=2, target=positioned.id,
                   before=(positioned.x, positioned.y),
                   after=(positioned.x + 457200, positioned.y)),
        ), out)

        text = diff(adversarial_deck, out).render()
        assert "1 change what the deck says" in text
        assert "1 change how it looks" in text

    def test_a_long_diff_is_truncated_with_a_count(self, adversarial_deck, tmp_path):
        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        result = diff(adversarial_deck, copy)
        from slide_wright.diff import ShapeDelta
        result.deltas = [
            ShapeDelta(slide=1, shape_id=str(i), kind="text", description=f"change {i}")
            for i in range(60)
        ]
        text = result.render(limit=10)
        assert "50 more difference(s)" in text


class TestAcceptsInspectedDecks:
    def test_takes_deckinfo_so_a_caller_need_not_parse_twice(self, adversarial_deck, tmp_path):
        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        assert not diff(inspect(adversarial_deck), inspect(copy)).changed


class TestFigureChanges:
    """The sharpest claim this product makes is not "content unchanged".

    It is *no figure changed*. Someone asking for a formatting pass on a
    pitchbook does not want reassurance about prose; they want to know the
    multiples are the ones they signed off.
    """

    def delta(self, kind, before, after):
        return ShapeDelta(slide=1, shape_id="1", kind=kind,
                          description="", before=before, after=after)

    def test_a_moved_number_is_a_figure_change(self):
        assert self.delta("text", "9.4x", "11.8x").changes_figures
        assert self.delta("table", "312", "340").changes_figures

    def test_presentation_is_never_a_figure_change(self):
        """A typeface carrying digits must not read as a moved number."""
        assert not self.delta("formatting", "Arial 10", "Arial 12").changes_figures

    def test_prose_edited_without_touching_a_number_is_not(self):
        assert not self.delta("text", "teh margin", "the margin").changes_figures
        assert not self.delta("text", "Revenue", "Revenue up").changes_figures

    def test_reordering_digits_counts(self):
        """Comparing sets would call this unchanged. It is a different deck."""
        assert self.delta("text", "9.4 and 4.9", "4.9 and 9.4").changes_figures

    def test_it_errs_toward_flagging(self):
        """This claim is a negative being proved.

        A false alarm costs a second look; a missed one costs the guarantee. So
        a label that merely contains a digit is flagged rather than parsed for
        whether it is really a quantity.
        """
        assert self.delta("text", "Q3 2026", "Q4 2026").changes_figures

    def test_the_deck_reports_them_separately(self, adversarial_deck, table, tmp_path):
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table.id}/r1/c1", before="9.4x", after="11.8x"))
        apply_changes(adversarial_deck, cs, out)

        result = diff(adversarial_deck, out)
        assert result.figure_deltas, "a changed multiple is a figure change"
        assert all(d.is_content for d in result.figure_deltas)


class TestFormattingIsComparedEvenWhenTheRunsMoved:
    """Same words, different emphasis — the case the diff used to call identical.

    Runs are an implementation detail of OOXML that PowerPoint rearranges freely:
    typing into a line merges runs, applying a style splits them. Comparing
    formatting only when both sides happen to have the same number of runs meant
    the single most misleading answer this product can give — "no structural
    differences" — for two decks that read the same and look different.
    """

    def _deck(self, path, pieces):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text, bold in pieces:
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(18)
        prs.save(str(path))
        return path

    def test_a_lost_emphasis_is_reported_when_runs_were_merged(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx", (("Total ", False), ("42", True)))
        after = self._deck(tmp_path / "b.pptx", (("Total 42", False),))

        result = diff(before, after)
        assert result.changed, "two decks that look different are not identical"
        assert len(result.deltas) == 1
        delta = result.deltas[0]
        assert delta.kind == "formatting"
        assert not delta.is_content, "emphasis is presentation, not content"
        assert delta.summary == "bold True -> False"
        assert "'42'" in delta.description, "it must name the text that changed"

    def test_equal_run_counts_do_not_make_index_pairing_safe(self, tmp_path):
        """One boundary moved by a character. Both sides have two runs."""
        before = self._deck(tmp_path / "a.pptx", (("Total ", False), ("42", True)))
        after = self._deck(tmp_path / "b.pptx", (("Total", False), (" 42", True)))

        deltas = diff(before, after).deltas
        assert [d.summary for d in deltas] == ["bold False -> True"], (
            "the space did change weight; pairing run 1 with run 1 misses it"
        )
        assert "' '" in deltas[0].description

    def test_the_same_sentence_split_two_ways_is_not_a_difference(self, tmp_path):
        before = self._deck(tmp_path / "c.pptx", (("Tot", False), ("al 42", False)))
        after = self._deck(tmp_path / "d.pptx", (("Total ", False), ("42", False)))

        assert not diff(before, after).changed, (
            "identical formatting at every character is no difference"
        )

    def test_one_span_is_reported_once_not_once_per_character(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx", (("Revenue and margin", True),))
        after = self._deck(
            tmp_path / "b.pptx", (("Revenue", True), (" and margin", False))
        )

        deltas = diff(before, after).deltas
        assert [d.summary for d in deltas] == ["bold True -> False"]
        assert "' and margin'" in deltas[0].description


class TestALostLinkIsReported:
    """A hyperlink is carried by a run, so an edit can take one off the slide.

    The `a:hlinkClick` and the relationship both survive -- the run keeps its
    `rPr`, the .rels file keeps its target -- and there is no longer any text
    carrying the link. Part counts are unchanged, the package is well formed,
    and the deck has a dead link.

    Nothing in this product could see that until now: `inspect` did not record a
    run's link, so the diff had nothing to compare and the audit had nothing to
    check. The only line printed was the text delta, which says nothing about
    where the deck used to point.
    """

    def _linked(self, path):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text in ("See ", "methodology", " for details"):
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(18)
            if text == "methodology":
                run.hyperlink.address = "https://example.com/methodology"
        prs.save(str(path))
        return path

    def _shape(self, deck):
        return next(s for s in inspect(deck).slides[0].shapes if s.text.strip())

    def test_inspect_records_where_a_run_points(self, tmp_path):
        src = self._linked(tmp_path / "a.pptx")
        assert [r.link for r in self._shape(src).runs] == [
            None, "https://example.com/methodology", None
        ], "a run's link has to be readable before anything can compare it"

    def test_an_edit_that_drops_a_link_says_so(self, tmp_path):
        from slide_wright.apply import apply_changes
        from slide_wright.changeset import Change, ChangeSet, Op

        src = self._linked(tmp_path / "a.pptx")
        shape = self._shape(src)
        out = tmp_path / "b.pptx"
        cs = ChangeSet(deck=str(src))
        cs.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                      before="See methodology", after="See the appendix"))
        cs.approve_all()
        assert not apply_changes(src, cs, out).failed

        deltas = diff(src, out).deltas
        lost = [d for d in deltas if d.kind == "link"]
        assert len(lost) == 1, [d.description for d in deltas]
        assert "https://example.com/methodology" in lost[0].description
        assert "removed" in lost[0].description

    def test_a_dropped_link_counts_as_content(self, tmp_path):
        """It is what the deck points at, and a reader notices losing one."""
        from slide_wright.apply import apply_changes
        from slide_wright.changeset import Change, ChangeSet, Op

        src = self._linked(tmp_path / "a.pptx")
        shape = self._shape(src)
        out = tmp_path / "b.pptx"
        cs = ChangeSet(deck=str(src))
        cs.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                      before="See methodology", after="See the appendix"))
        cs.approve_all()
        apply_changes(src, cs, out)

        result = diff(src, out)
        assert all(d.is_content for d in result.deltas if d.kind == "link")

    def test_a_link_left_alone_is_not_reported(self, tmp_path):
        src = self._linked(tmp_path / "a.pptx")
        again = self._linked(tmp_path / "b.pptx")
        assert not [d for d in diff(src, again).deltas if d.kind == "link"]


class TestFormattingAReaderSeesAndTheDiffDidNot:
    """Bold and italic were compared. Underline, strikethrough and baseline were not.

    Three attributes on the same element as the two that *were* read, each one
    visible at a glance: "Confidential draft" losing its underline, a figure
    gaining a strikethrough, a footnote marker dropping out of superscript. Every
    one of them came back as **no structural differences**.

    Measured across the 26 real fixtures: 50 underlined runs, 24 with a baseline
    and 6 struck through — 80 runs carrying formatting nothing here could see,
    and superscript in 5 of the 26 decks.
    """

    def _deck(self, path, **rpr):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        run = frame.paragraphs[0].add_run()
        run.text = "Confidential draft"
        run.font.size = Pt(18)
        for name, value in rpr.items():
            run.font._rPr.set(name, value)
        prs.save(str(path))
        return path

    def test_a_lost_underline_is_reported(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx", u="sng")
        after = self._deck(tmp_path / "b.pptx")
        deltas = diff(before, after).deltas
        assert [d.summary for d in deltas] == ["underline 'sng' -> None"]

    def test_a_double_underline_becoming_single_is_reported(self, tmp_path):
        """Flattening these to booleans would have hidden this one."""
        before = self._deck(tmp_path / "a.pptx", u="dbl")
        after = self._deck(tmp_path / "b.pptx", u="sng")
        assert [d.summary for d in diff(before, after).deltas] == [
            "underline 'dbl' -> 'sng'"
        ]

    def test_a_strikethrough_is_reported(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx")
        after = self._deck(tmp_path / "b.pptx", strike="sngStrike")
        assert [d.summary for d in diff(before, after).deltas] == [
            "strikethrough None -> 'sngStrike'"
        ]

    def test_a_lost_superscript_is_reported(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx", baseline="30000")
        after = self._deck(tmp_path / "b.pptx")
        assert [d.summary for d in diff(before, after).deltas] == [
            "baseline 30000 -> None"
        ]

    def test_saying_off_and_saying_nothing_are_not_a_difference(self, tmp_path):
        """`u="none"` and no `u` at all render identically, so they must compare
        identically. Reporting that gap would be a difference nobody can see."""
        before = self._deck(tmp_path / "a.pptx", u="none", strike="noStrike", baseline="0")
        after = self._deck(tmp_path / "b.pptx")
        assert not diff(before, after).changed

    def test_these_are_presentation_not_content(self, tmp_path):
        before = self._deck(tmp_path / "a.pptx", u="sng")
        after = self._deck(tmp_path / "b.pptx")
        assert all(not d.is_content for d in diff(before, after).deltas)


class TestEmphasisLostToATextEdit:
    """Styling a shape used before an edit and does not use after it.

    Run formatting is deliberately not compared when the text changed: the runs
    have been re-described by the text delta, and reporting every field of every
    rewritten run buries the line the reviewer needs. That holds for *changes*
    and not for *disappearances*.

    It matters because replacing a shape's text is the only way the workspace
    edits words — the contract names an object and its new full text — so every
    intra-shape emphasis collapses into one run on every edit. Changing FY25 to
    FY26 on "Revenue grew **15%** in FY25" took the bold off the figure and the
    diff said `text ...FY2[5 -> 6]`.
    """

    def _deck(self, path, runs):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(1)
        ).text_frame
        for text, bold, size in runs:
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(size)
        prs.save(str(path))
        return path

    def _rewrite(self, src, tmp_path, after):
        from slide_wright.apply import apply_changes
        from slide_wright.changeset import Change, ChangeSet, Op

        shape = next(s for s in inspect(src).slides[0].shapes if s.text.strip())
        out = tmp_path / "after.pptx"
        cs = ChangeSet(deck=str(src))
        cs.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                      before=shape.text, after=after))
        cs.approve_all()
        assert not apply_changes(src, cs, out).failed
        return diff(src, out)

    def test_a_flattened_emphasis_is_reported(self, tmp_path):
        src = self._deck(tmp_path / "a.pptx", [
            ("Revenue grew ", False, 18), ("15%", True, 24), (" in FY25", False, 18),
        ])
        summaries = [d.summary for d in self._rewrite(src, tmp_path, "Revenue grew 15% in FY26").deltas]
        assert "bold True no longer used in this text" in summaries
        assert "size 24.0 no longer used in this text" in summaries

    def test_a_uniform_shape_reports_nothing_extra(self, tmp_path):
        """Nothing to lose: there was one style and there still is."""
        src = self._deck(tmp_path / "a.pptx", [("Revenue grew in FY25", False, 18)])
        kinds = [d.kind for d in self._rewrite(src, tmp_path, "Revenue grew in FY26").deltas]
        assert kinds == ["text"]

    def test_emphasis_that_survives_is_not_reported(self, tmp_path):
        """The rewrite keeps a bold run, so bold is still in use."""
        from slide_wright.apply import apply_changes
        from slide_wright.changeset import Change, ChangeSet, Op

        src = self._deck(tmp_path / "a.pptx", [
            ("Revenue grew ", False, 18), ("15%", True, 18),
        ])
        shape = next(s for s in inspect(src).slides[0].shapes if s.text.strip())
        out = tmp_path / "b.pptx"
        cs = ChangeSet(deck=str(src))
        # A narrow edit inside the first run: the bold run is untouched.
        cs.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                      before="Revenue grew", after="Revenue rose"))
        cs.approve_all()
        assert not apply_changes(src, cs, out).failed
        assert not [d for d in diff(src, out).deltas if "no longer used" in d.summary]

    def test_it_is_presentation_not_content(self, tmp_path):
        src = self._deck(tmp_path / "a.pptx", [
            ("Revenue grew ", False, 18), ("15%", True, 24),
        ])
        lost = [d for d in self._rewrite(src, tmp_path, "Revenue grew 20%").deltas
                if "no longer used" in d.summary]
        assert lost and all(not d.is_content for d in lost)


class TestHowALineSitsIsCompared:
    """Bullets, indent, alignment and line spacing — none of them were read.

    The diff compared eight run attributes and nothing at all about the
    paragraph a run sits in, so a deck whose bullets had been taken off, whose
    body had been centred, or whose lines had been tightened came back *"no
    structural differences"*. Every verifier-side lock is a question put to this
    module, so a property it cannot compute is a property no lock can protect --
    including `formatting`, whose whole sentence is *leave my styling exactly as
    it is*.

    Measured across the corpus before building it: 1,635 explicitly aligned
    paragraphs in 10 of 26 decks, 1,137 declaring a bullet in 8, 707 setting
    their own line spacing, 212 indented.
    """

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def _bulleted_part(self, deck) -> str:
        import re
        import zipfile

        with zipfile.ZipFile(deck) as z:
            for name in sorted(z.namelist()):
                if re.match(r"^ppt/slides/slide\d+\.xml$", name) and b"Margin held" in z.read(name):
                    return name
        raise AssertionError("the corpus has no bulleted body to ask about")

    def _variant(self, deck, out, mutate):
        """The same deck with one slide part rewritten. Nothing else moves."""
        import zipfile

        from lxml import etree

        part = self._bulleted_part(deck)
        with zipfile.ZipFile(deck) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == part:
                    root = etree.fromstring(data)
                    paragraph = next(
                        para for para in root.iter(f"{self.A}p")
                        if para.findall(f".//{self.A}t")
                    )
                    mutate(paragraph, etree)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                          standalone=True)
                dst.writestr(item, data)
        return out

    def _pPr(self, paragraph, etree):
        existing = paragraph.find(f"{self.A}pPr")
        if existing is not None:
            return existing
        created = etree.Element(f"{self.A}pPr")
        paragraph.insert(0, created)
        return created

    def _summaries(self, deck, out):
        return [d.summary for d in diff(deck, out).deltas]

    def test_a_bullet_taken_off_is_reported(self, adversarial_deck, tmp_path):
        out = self._variant(
            adversarial_deck, tmp_path / "o.pptx",
            lambda para, etree: etree.SubElement(self._pPr(para, etree), f"{self.A}buNone"),
        )
        assert "bullet inherited -> none" in self._summaries(adversarial_deck, out)

    def test_a_line_centred_is_reported(self, adversarial_deck, tmp_path):
        out = self._variant(
            adversarial_deck, tmp_path / "o.pptx",
            lambda para, etree: self._pPr(para, etree).set("algn", "ctr"),
        )
        assert "alignment inherited -> ctr" in self._summaries(adversarial_deck, out)

    def test_an_indent_level_is_reported(self, adversarial_deck, tmp_path):
        out = self._variant(
            adversarial_deck, tmp_path / "o.pptx",
            lambda para, etree: self._pPr(para, etree).set("lvl", "2"),
        )
        assert "indent level 0 -> 2" in self._summaries(adversarial_deck, out)

    def test_tighter_line_spacing_is_reported(self, adversarial_deck, tmp_path):
        def tighten(para, etree):
            spacing = etree.SubElement(self._pPr(para, etree), f"{self.A}lnSpc")
            etree.SubElement(spacing, f"{self.A}spcPct").set("val", "80000")

        out = self._variant(adversarial_deck, tmp_path / "o.pptx", tighten)
        assert "line spacing inherited -> 80%" in self._summaries(adversarial_deck, out)

    def test_it_is_how_the_deck_looks_not_what_it_says(self, adversarial_deck, tmp_path):
        out = self._variant(
            adversarial_deck, tmp_path / "o.pptx",
            lambda para, etree: self._pPr(para, etree).set("algn", "ctr"),
        )
        result = diff(adversarial_deck, out)
        assert result.deltas and not result.content_deltas, (
            "moving a line is presentation; it changes nothing the deck says"
        )

    def test_the_line_is_named(self, adversarial_deck, tmp_path):
        out = self._variant(
            adversarial_deck, tmp_path / "o.pptx",
            lambda para, etree: self._pPr(para, etree).set("algn", "ctr"),
        )
        delta = next(d for d in diff(adversarial_deck, out).deltas if "alignment" in d.summary)
        assert " line 1 " in delta.description, (
            "a shape with four lines needs to say which one moved"
        )

    def test_properties_are_not_paired_when_the_lines_changed(
        self, adversarial_deck, tmp_path
    ):
        """Index pairing is only sound while the counts agree.

        When they do not, the count delta has already said so -- which is the
        difference from the run comparison, where an early return on a count
        mismatch produced *"no structural differences"* over a deck whose figure
        had lost its bold.
        """
        def drop(para, etree):
            self._pPr(para, etree).set("algn", "ctr")
            para.getparent().remove(para)

        out = self._variant(adversarial_deck, tmp_path / "o.pptx", drop)
        summaries = self._summaries(adversarial_deck, out)
        assert any("text lines" in x for x in summaries), "the lost line must be reported"
        assert not any("alignment" in x for x in summaries), (
            "lines cannot be paired by index once one of them has gone"
        )


class TestInheritedIsNotAbsent:
    """A line with no properties of its own takes them from its layout."""

    def test_a_paragraph_with_no_properties_reads_as_inherited(self):
        from lxml import etree

        from slide_wright.inspect import _paragraph

        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        para = etree.fromstring(
            f'<a:p xmlns:a="{A[1:-1]}"><a:r><a:t>Alpha</a:t></a:r></a:p>'
        )
        info = _paragraph(para)
        assert (info.level, info.alignment, info.bullet, info.line_spacing) == (
            0, None, None, None,
        )

    def test_an_off_state_is_a_value_not_an_absence(self):
        """`buNone` is a decision. It is not the same as saying nothing."""
        from lxml import etree

        from slide_wright.inspect import _paragraph

        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        para = etree.fromstring(
            f'<a:p xmlns:a="{A[1:-1]}"><a:pPr><a:buNone/></a:pPr>'
            "<a:r><a:t>Alpha</a:t></a:r></a:p>"
        )
        assert _paragraph(para).bullet == "none"


class TestCapitalsAreReadAndCompared:
    """A run that says one thing on the slide and another in the XML.

    `cap="all"` is the ninth attribute on the same element as `bold` and
    `italic`, and like the three found on 8 Sep it was not read. It is the one
    with the loudest effect: a section header reading DIVIDER is stored as
    "divider", so a text edit that moves its words into a neighbouring run
    changes what a reader sees while changing no character of the text.

    Measured first, and the measurement is why the corpus had to change: every
    one of the 363 `cap` attributes across the 26 real decks is `cap="none"`,
    the off state. A comparison written against them could never fire, which
    reads as coverage and is not -- so the adversarial deck now carries a run
    that is genuinely ALL CAPS.
    """

    def _caps_run(self, deck):
        return next(
            (sl.number, sh)
            for sl in inspect(deck).slides
            for sh in sl.shapes
            if any(r.caps for r in sh.runs)
        )

    def test_the_corpus_carries_a_run_that_is_actually_capitalised(
        self, adversarial_deck
    ):
        _, shape = self._caps_run(adversarial_deck)
        assert [r.caps for r in shape.runs] == [None, "all"], (
            "the fixture is the whole point; without it nothing below can fail"
        )

    def test_an_edit_that_takes_the_capitals_off_says_so(
        self, adversarial_deck, tmp_path
    ):
        number, shape = self._caps_run(adversarial_deck)
        out = tmp_path / "o.pptx"
        apply_changes(adversarial_deck, approved(adversarial_deck, Change(
            id="c", op=Op.SET_TEXT, slide=number, target=shape.id,
            before=shape.text, after="Section header",
        )), out)

        summaries = [d.summary for d in diff(adversarial_deck, out).deltas]
        assert any("capitals" in x for x in summaries), (
            f"the deck lost its ALL CAPS and the diff said {summaries}"
        )

    def test_an_off_state_is_not_a_difference(self):
        """`cap="none"` renders identically to no attribute at all.

        363 runs in the corpus carry it. Reporting the two apart would produce a
        difference nobody can see, which is the reasoning already written down
        for underline, strikethrough and baseline.
        """
        from lxml import etree

        from slide_wright.inspect import _off

        assert _off("none", "none") is None
        assert _off(None, "none") is None
        assert _off("all", "none") == "all"


def rewrite_notes(deck, out, replace: bytes, with_: bytes):
    """The same deck with one word changed in one notes part. Nothing else moves.

    Built by hand rather than by the applier because nothing in this engine
    writes a notes part -- which is exactly why the diff could not see one
    change. A second line of defence never fires on anything the first line
    produces.
    """
    import re
    import zipfile

    target = None
    with zipfile.ZipFile(deck) as z:
        for name in sorted(z.namelist()):
            if re.match(r"^ppt/notesSlides/notesSlide\d+\.xml$", name) and replace in z.read(name):
                target = name
                break
        assert target, f"no notes part in {deck.name} contains {replace!r}"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == target:
                    data = data.replace(replace, with_, 1)
                dst.writestr(item, data)
    return out


class TestSpeakerNotesAreCompared:
    """A third of a real deck's words lived where no comparison could reach.

    Notes are in their own part, so they survive an edit and `verify` compares
    them byte for byte. Meaning was the gap: change a word of the presenter's
    script and this module said *"No structural differences."* It reaches
    further than a report, because every verifier-side lock is a question put
    to the diff.

    Measured before building: 43 notes slides across the corpus carry 3,212
    words, and on `nasa-bhutan-water` the script is 2,708 words against 983 on
    the slides.
    """

    def test_a_word_changed_in_the_script_is_reported(self, adversarial_deck, tmp_path):
        out = rewrite_notes(
            adversarial_deck, tmp_path / "o.pptx", b"presenter", b"narrator"
        )
        result = diff(adversarial_deck, out)
        assert result.changed, "changing the script must not read as no difference"
        assert any(d.kind == "notes" for d in result.deltas)

    def test_it_is_what_the_deck_says(self, adversarial_deck, tmp_path):
        out = rewrite_notes(
            adversarial_deck, tmp_path / "o.pptx", b"presenter", b"narrator"
        )
        delta = next(d for d in diff(adversarial_deck, out).deltas if d.kind == "notes")
        assert delta.is_content, "a presenter's script is words somebody wrote"
        assert "speaker notes" in delta.description

    def test_it_names_the_slide_and_claims_no_shape(self, adversarial_deck, tmp_path):
        out = rewrite_notes(
            adversarial_deck, tmp_path / "o.pptx", b"presenter", b"narrator"
        )
        delta = next(d for d in diff(adversarial_deck, out).deltas if d.kind == "notes")
        assert delta.slide > 0
        assert delta.shape_id == "", "notes belong to the slide, not to an object on it"

    def test_a_figure_changed_in_the_script_is_flagged(self, adversarial_deck, tmp_path):
        """`changes_figures` is the sharpest claim here, and it must cover notes.

        A number in the script is a number somebody will read aloud.
        """
        out = rewrite_notes(
            adversarial_deck, tmp_path / "o.pptx", b"presenter", b"2026 presenter"
        )
        assert diff(adversarial_deck, out).figure_deltas

    def test_an_identical_deck_still_reports_nothing(self, adversarial_deck, tmp_path):
        import shutil

        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        assert not diff(adversarial_deck, copy).changed
