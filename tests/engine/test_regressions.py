"""Regressions.

Every bug found in real use gets a permanent test here, named for what it did
rather than for the code that was wrong.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from slide_wright.apply import apply_changes
from slide_wright.changeset import Change, ChangeSet, Op
from slide_wright.inspect import inspect
from slide_wright.session import Session, SessionError


def multi_paragraph_deck(path):
    """A shape whose text spans several paragraphs and several runs.

    This is the ordinary shape of a real content slide, and it is what broke.
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(4))
    frame = box.text_frame
    frame.text = "First paragraph of the body"
    second = frame.add_paragraph()
    run_a = second.add_run()
    run_a.text = "Second paragraph "
    run_a.font.bold = True          # forces a run split
    run_b = second.add_run()
    run_b.text = "continues here"
    run_b.font.size = Pt(14)
    prs.save(str(path))
    return path


class TestSilentNoOpApply:
    """apply() reported "VERIFIED" while applying nothing.

    Found by the Phase 1 exit check: 8 of 12 real decks produced an unchanged
    file and a report claiming success. A no-op that looks like success is the
    worst failure this product can have, because nobody investigates a pass.
    """

    def test_session_refuses_to_commit_when_nothing_applied(self, adversarial_deck, tmp_path):
        session = Session.open(adversarial_deck, workspace=tmp_path / "ws")
        changeset = session.propose("edit text that is not there")
        changeset.add(Change(
            id="c1", op=Op.SET_TEXT, slide=1, target="2",
            before="this string is nowhere in the deck", after="replacement",
        ))
        changeset.approve_all()

        with pytest.raises(SessionError, match="could not be applied"):
            session.apply()

    def test_session_stays_on_the_previous_version_after_a_failed_apply(
        self, adversarial_deck, tmp_path
    ):
        session = Session.open(adversarial_deck, workspace=tmp_path / "ws")
        changeset = session.propose("bad edit")
        changeset.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target="2",
                             before="absent text", after="x"))
        changeset.approve_all()
        with pytest.raises(SessionError):
            session.apply()
        assert session.current.number == 0

    def test_partial_failure_does_not_pass_as_success(self, adversarial_deck, tmp_path):
        """One good change and one impossible change must not read as success."""
        session = Session.open(adversarial_deck, workspace=tmp_path / "ws")
        deck = session.deck()
        title = deck.slide(1).shapes[0]

        changeset = session.propose("one good, one impossible")
        changeset.add(Change(id="good", op=Op.SET_TEXT, slide=1, target=title.id,
                             before=title.text, after="New Title"))
        changeset.add(Change(id="bad", op=Op.SET_TEXT, slide=1, target="2",
                             before="absent text", after="x"))
        changeset.approve_all()

        with pytest.raises(SessionError, match="could not be applied"):
            session.apply()


class TestTextSpanningParagraphs:
    """set_text failed whenever the target text crossed a paragraph boundary.

    `ShapeInfo.text` joins every run in the shape, so a caller passing that
    value back could describe a span no single paragraph contained. The applier
    only searched within paragraphs, so it silently matched nothing.

    The join now puts a newline between paragraphs, so the string the caller
    reads carries the shape's line breaks. The property under test is unchanged
    and slightly stronger: whatever `ShapeInfo.text` returns, passing it back
    as `before` has to find the whole shape.
    """

    def test_replaces_text_spanning_paragraphs_and_runs(self, tmp_path):
        deck_path = multi_paragraph_deck(tmp_path / "multi.pptx")
        shape = next(s for s in inspect(deck_path).all_shapes() if s.has_text)
        assert "\n" in shape.text, "paragraphs are separated, not welded"

        changeset = ChangeSet(deck=str(deck_path))
        changeset.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                             before=shape.text, after="Replaced entirely"))
        changeset.approve_all()

        out = tmp_path / "out.pptx"
        result = apply_changes(deck_path, changeset, out)

        assert result.ok, f"apply failed: {result.failed}"
        assert "Replaced entirely" in inspect(out).slide(1).text

    def test_replaces_text_spanning_runs_within_one_paragraph(self, tmp_path):
        deck_path = multi_paragraph_deck(tmp_path / "multi.pptx")
        changeset = ChangeSet(deck=str(deck_path))
        shape = next(s for s in inspect(deck_path).all_shapes() if s.has_text)
        changeset.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
                             before="Second paragraph continues here",
                             after="Rewritten sentence"))
        changeset.approve_all()

        out = tmp_path / "out.pptx"
        assert apply_changes(deck_path, changeset, out).ok
        assert "Rewritten sentence" in inspect(out).slide(1).text

    def test_single_run_match_is_still_preferred(self, adversarial_deck, tmp_path):
        """The narrow path must not regress now that broader ones exist."""
        import difflib
        import zipfile

        deck = inspect(adversarial_deck)
        table = next(s for s in deck.all_shapes() if s.kind == "table")
        changeset = ChangeSet(deck=str(adversarial_deck))
        changeset.add(Change(id="c1", op=Op.SET_TABLE_CELL, slide=3,
                             target=f"{table.id}/r1/c1", before="9.4x", after="11.8x"))
        changeset.approve_all()
        out = tmp_path / "out.pptx"
        apply_changes(adversarial_deck, changeset, out)

        def body(p):
            raw = zipfile.ZipFile(p).read("ppt/slides/slide3.xml").decode()
            return raw[raw.index("<p:sld"):]

        matcher = difflib.SequenceMatcher(None, body(adversarial_deck), body(out))
        edits = [op for op in matcher.get_opcodes() if op[0] != "equal"]
        assert len(edits) <= 3, "narrow single-run replacement regressed"


class TestImageOnlyDecks:
    """A deck with no text is a legitimate input, not an error."""

    def test_inspecting_an_image_only_deck_works(self, tmp_path):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])  # blank
        path = tmp_path / "blank.pptx"
        prs.save(str(path))

        deck = inspect(path)
        assert deck.slide_count == 1
        assert deck.slides[0].word_count == 0
        assert deck.slides[0].title is None
