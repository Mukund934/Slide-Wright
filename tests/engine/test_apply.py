"""In-place application of approved changes.

The property under test is narrowness: an edit must touch the target and
nothing else. Slide-level preservation is table stakes here — these tests
assert preservation *within* the edited slide, down to the character.
"""

from __future__ import annotations

import difflib
import re
import zipfile
from dataclasses import fields
from pathlib import Path

import pytest

from slide_wright.apply import ApplyError, apply_changes
from slide_wright.changeset import Change, ChangeSet, Op, Status
from slide_wright.fidelity import compare
from slide_wright.inspect import TextRun, inspect
from slide_wright.package import Package


def slide_body(path, part: str) -> str:
    raw = zipfile.ZipFile(path).read(part).decode("utf-8")
    return raw[raw.index("<p:sld"):]


def approved(deck, *changes: Change) -> ChangeSet:
    cs = ChangeSet(deck=str(deck))
    for c in changes:
        cs.add(c)
    cs.approve_all()
    return cs


@pytest.fixture
def table_shape(adversarial_deck):
    return next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")


class TestNarrowness:
    def test_a_cell_edit_changes_two_characters(self, adversarial_deck, table_shape, tmp_path):
        """The strongest form of the promise, measured rather than asserted."""
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        apply_changes(adversarial_deck, cs, out)

        before = slide_body(adversarial_deck, "ppt/slides/slide3.xml")
        after = slide_body(out, "ppt/slides/slide3.xml")
        matcher = difflib.SequenceMatcher(None, before, after)
        edits = [op for op in matcher.get_opcodes() if op[0] != "equal"]
        preserved = sum(b.size for b in matcher.get_matching_blocks()) / len(before)

        assert len(edits) <= 3, f"expected a surgical edit, got {edits}"
        assert preserved > 0.99, f"only {preserved:.2%} of the edited slide preserved"

    def test_only_the_target_slide_part_is_touched(self, adversarial_deck, table_shape, tmp_path):
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        result = apply_changes(adversarial_deck, cs, out)
        assert result.touched_parts == {"ppt/slides/slide3.xml"}
        assert [d.name for d in compare(adversarial_deck, out).changed] == [
            "ppt/slides/slide3.xml"
        ]

    def test_run_count_is_unchanged_by_a_text_edit(self, adversarial_deck, tmp_path):
        deck_info = inspect(adversarial_deck)
        title = deck_info.slide(1).shapes[0]
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target=title.id,
            before=title.text, after="Renamed Corpus",
        ))
        apply_changes(adversarial_deck, cs, out)
        rep = compare(adversarial_deck, out)
        assert rep.output_census.text_runs == rep.source_census.text_runs


def movable(deck, slide_number: int):
    """A plain shape with geometry of its own, on the given slide.

    Three things disqualify a shape here and each is a real engine rule, not a
    fixture convenience: a placeholder is positioned by its layout and the
    applier refuses to move it; a chart is read-only (ADR-0009); and SmartArt is
    read-only (ADR-0007). Picking "the first shape with an x" used to land on
    whichever of those came first and made the test assert the wrong refusal.
    """
    return next(
        s for s in inspect(deck).slide(slide_number).shapes
        if s.kind == "shape" and s.x is not None and not s.geometry_inherited
    )


class TestOperations:
    def test_set_text_writes_the_new_value(self, adversarial_deck, tmp_path):
        d = inspect(adversarial_deck)
        title = d.slide(1).shapes[0]
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target=title.id,
            before=title.text, after="Renamed Corpus",
        ))
        apply_changes(adversarial_deck, cs, out)
        assert inspect(out).slide(1).title == "Renamed Corpus"

    def test_set_table_cell_writes_the_new_value(self, adversarial_deck, table_shape, tmp_path):
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        apply_changes(adversarial_deck, cs, out)
        assert "11.8x" in inspect(out).slide(3).text
        assert "9.4x" not in inspect(out).slide(3).text

    def test_move_updates_offset(self, adversarial_deck, tmp_path):
        shape = movable(adversarial_deck, 6)
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.MOVE, slide=6, target=shape.id,
            before=(shape.x, shape.y), after=(shape.x + 100000, shape.y),
        ))
        apply_changes(adversarial_deck, cs, out)
        moved = next(s for s in inspect(out).slide(6).shapes if s.id == shape.id)
        assert moved.x == shape.x + 100000

    def test_resize_updates_extent(self, adversarial_deck, tmp_path):
        shape = movable(adversarial_deck, 6)
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.RESIZE, slide=6, target=shape.id,
            before=(shape.cx, shape.cy), after=(shape.cx // 2, shape.cy),
        ))
        apply_changes(adversarial_deck, cs, out)
        resized = next(s for s in inspect(out).slide(6).shapes if s.id == shape.id)
        assert resized.cx == shape.cx // 2


class TestRefusals:
    """Fail closed. Never approximate a change we cannot make precisely."""

    def test_unsupported_op_is_refused_not_approximated(self, adversarial_deck, tmp_path):
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.DELETE_SHAPE, slide=1, target="2",
        ))
        with pytest.raises(ApplyError, match="cannot perform"):
            apply_changes(adversarial_deck, cs, tmp_path / "o.pptx")

    def test_empty_change_set_is_refused(self, adversarial_deck, tmp_path):
        with pytest.raises(ApplyError, match="no approved changes"):
            apply_changes(adversarial_deck, ChangeSet(deck=str(adversarial_deck)),
                          tmp_path / "o.pptx")

    def test_rejected_changes_are_not_applied(self, adversarial_deck, tmp_path):
        d = inspect(adversarial_deck)
        title = d.slide(1).shapes[0]
        cs = ChangeSet(deck=str(adversarial_deck))
        cs.add(Change(id="c1", op=Op.SET_TEXT, slide=1, target=title.id,
                      before=title.text, after="Should Not Appear"))
        cs.add(Change(id="c2", op=Op.SET_TEXT, slide=1, target=title.id,
                      before=title.text, after="Should Appear"))
        cs.approve("c2")
        cs.reject("c1")
        out = tmp_path / "o.pptx"
        apply_changes(adversarial_deck, cs, out)
        assert inspect(out).slide(1).title == "Should Appear"

    def test_missing_slide_is_recorded_as_failure(self, adversarial_deck, tmp_path):
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=999, target="1", before="x", after="y",
        ))
        result = apply_changes(adversarial_deck, cs, tmp_path / "o.pptx")
        assert not result.ok
        assert "not found" in result.failed[0][1]

    def test_missing_target_is_recorded_as_failure(self, adversarial_deck, tmp_path):
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target="99999", before="x", after="y",
        ))
        result = apply_changes(adversarial_deck, cs, tmp_path / "o.pptx")
        assert not result.ok


class TestSourceIsImmutable:
    def test_the_source_file_is_never_modified(self, adversarial_deck, table_shape, tmp_path):
        original = adversarial_deck.read_bytes()
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        apply_changes(adversarial_deck, cs, tmp_path / "o.pptx")
        assert adversarial_deck.read_bytes() == original


class TestIntegrityPreserved:
    def test_native_objects_survive_an_edit(self, adversarial_deck, table_shape, tmp_path):
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        apply_changes(adversarial_deck, cs, out)
        rep = compare(adversarial_deck, out)
        assert not rep.native_losses
        assert not rep.rasterisation_suspected
        assert rep.structurally_intact

    def test_output_is_a_valid_package(self, adversarial_deck, table_shape, tmp_path):
        out = tmp_path / "o.pptx"
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=3,
            target=f"{table_shape.id}/r1/c1", before="9.4x", after="11.8x",
        ))
        apply_changes(adversarial_deck, cs, out)
        assert Package.open(out).slide_count == Package.open(adversarial_deck).slide_count


class TestFailuresExplainThemselves:
    """A refusal is only useful if it says what to do differently.

    `_apply_one` returned a single boolean, so "the shape is not on this slide"
    and "the shape is here but the edit does not apply to it" produced the same
    message: "target not found on slide". That sends a reader hunting for a
    shape id that is in fact present, and it fired on the common case of a run
    whose font size is inherited from the layout.
    """

    def _why(self, deck, change, tmp_path) -> str:
        cs = approved(deck, change)
        result = apply_changes(deck, cs, tmp_path / "out.pptx")
        assert result.failed, "expected this change to fail"
        return result.failed[0][1]

    def test_a_missing_shape_says_so_and_names_the_slide(self, adversarial_deck, tmp_path):
        why = self._why(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target="999",
            before="x", after="y"), tmp_path)
        assert "no shape with id '999'" in why and "slide 1" in why

    def test_wrong_before_text_names_the_text_not_the_shape(self, adversarial_deck, tmp_path):
        shape = next(s for s in inspect(adversarial_deck).slides[0].shapes if s.has_text)
        why = self._why(adversarial_deck, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
            before="text that is not in this deck", after="y"), tmp_path)
        assert "does not contain the text" in why
        assert "not found" not in why, "the shape was found; saying otherwise misleads"

    def test_an_inherited_font_size_is_explained_not_denied(self, adversarial_deck, tmp_path):
        """The commonest real case: a run with no explicit formatting override."""
        shape = movable(adversarial_deck, 2)
        why = self._why(adversarial_deck, Change(
            id="c1", op=Op.SET_FONT_SIZE, slide=2, target=shape.id,
            before=None, after=14), tmp_path)
        assert "inherited from the layout" in why
        assert "not found" not in why

    def test_a_layout_placed_shape_explains_why_it_cannot_move(
        self, adversarial_deck, tmp_path
    ):
        # Layout-placed is now visible as such rather than inferred from a
        # missing x: inspect resolves inherited geometry, so the shape has
        # coordinates and still has no position of its own.
        shape = next(s for s in inspect(adversarial_deck).slides[0].shapes
                     if s.geometry_inherited)
        why = self._why(adversarial_deck, Change(
            id="c1", op=Op.MOVE, slide=1, target=shape.id,
            before=(0, 0), after=(100, 100)), tmp_path)
        assert "placed by the layout" in why
        assert "not found" not in why


class TestRunAddressingMatchesInspect:
    """`<shape>/run/<index>` must mean the same run to both halves of the system.

    The index in a target is produced by reading the deck through `inspect`,
    which skips runs carrying no text. The applier used to enumerate every run,
    so a shape with two empty runs made inspect see fifteen where the applier
    saw seventeen — and a change addressed at run 13 was written to a different
    run entirely, while reporting success.

    Every check passed: content untouched, no native lost, fidelity fine. Only
    running the pass twice caught it, because the run that should have changed
    never did and kept being proposed again.
    """

    def _shape_element(self, deck, slide_part: str, shape_id: str):
        from lxml import etree

        from slide_wright.apply import NS
        from slide_wright.package import Package

        root = etree.fromstring(Package.open(deck).read(slide_part))
        for el in root.iter():
            if etree.QName(el).localname == "cNvPr" and el.get("id") == shape_id:
                return el.getparent().getparent()
        return None

    def test_the_applier_enumerates_what_inspect_enumerates(self, adversarial_deck):
        from slide_wright.apply import _addressable_runs

        deck = inspect(adversarial_deck)
        for slide in deck.slides:
            for shape in slide.shapes:
                if not shape.runs:
                    continue
                element = self._shape_element(adversarial_deck, slide.part_name, shape.id)
                if element is None:
                    continue
                assert len(_addressable_runs(element)) == len(shape.runs), (
                    f"slide {slide.number} shape {shape.id}: inspect sees "
                    f"{len(shape.runs)} runs, the applier sees "
                    f"{len(_addressable_runs(element))}"
                )

    def test_an_empty_run_is_not_addressable(self):
        from lxml import etree

        from slide_wright.apply import _addressable_runs

        xml = (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:p>'
            '<a:r><a:rPr/><a:t>first</a:t></a:r>'
            '<a:r><a:rPr/><a:t></a:t></a:r>'
            '<a:r><a:rPr/><a:t>second</a:t></a:r>'
            '</a:p></p:sp>'
        )
        runs = _addressable_runs(etree.fromstring(xml))
        assert len(runs) == 2, "an empty run must not consume an index"
        assert runs[1].find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        ).text == "second"


class TestMovingIsNotEditing:
    """The content guards were refusing an operation they have no argument against.

    A chart and a diagram each keep the same information twice — cached series
    against an embedded workbook, a diagram model against its drawing cache —
    and both guards exist because an edit that updates one copy and not the
    other produces a deck whose picture disagrees with its own data. That is
    the whole justification, and it is about content.

    Neither cache records where on the slide the frame sits. Refusing a move on
    those grounds cost more than it sounds: the alignment pass emits `MOVE` and
    nothing else, so one chart anywhere in a deck made `tidy` refuse the entire
    change set. Which is the demo.
    """

    CHART = (2, "3")   # slide, shape id — a native chart in the fixture
    TABLE = (3, "3")   # a native table

    def _move(self, deck, slide, target, tmp_path, name="moved.pptx"):
        pkg_shape = next(s for s in inspect(deck).slides[slide - 1].shapes
                         if s.id == target)
        cs = approved(deck, Change(
            id="m1", op=Op.MOVE, slide=slide, target=target,
            before=(pkg_shape.x, pkg_shape.y),
            after=(pkg_shape.x + 12700, pkg_shape.y),
        ))
        return apply_changes(deck, cs, tmp_path / name), pkg_shape

    def test_a_chart_frame_can_be_moved(self, adversarial_deck, tmp_path):
        result, _ = self._move(adversarial_deck, *self.CHART, tmp_path)
        assert result.ok, result.failed

    def test_the_move_actually_lands(self, adversarial_deck, tmp_path):
        """A graphicFrame states its box as `p:xfrm`, not `a:xfrm`.

        Looking only for the DrawingML form did not raise. It returned "shape 3
        has no position of its own; it is placed by the layout" — untrue of a
        graphicFrame, and it sends a reviewer looking for a layout that does not
        place it.
        """
        result, before = self._move(adversarial_deck, *self.CHART, tmp_path)
        after = next(s for s in inspect(result.output).slides[1].shapes if s.id == "3")
        assert after.x == before.x + 12700
        assert after.y == before.y

    def test_a_table_frame_can_be_moved_too(self, adversarial_deck, tmp_path):
        result, _ = self._move(adversarial_deck, *self.TABLE, tmp_path, "table.pptx")
        assert result.ok, result.failed

    def test_moving_a_chart_leaves_the_chart_part_untouched(self, adversarial_deck, tmp_path):
        """The reason the guard has no argument here, asserted rather than argued."""
        result, _ = self._move(adversarial_deck, *self.CHART, tmp_path, "parts.pptx")
        rep = compare(adversarial_deck, result.output)
        assert not [d.name for d in rep.changed if d.name.startswith("ppt/charts/")]
        assert not [d.name for d in rep.changed if d.name.startswith("ppt/embeddings/")]

    def test_editing_a_chart_is_still_refused(self, adversarial_deck, tmp_path):
        """The scoping must not become a way around the guard."""
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_FONT, slide=self.CHART[0], target=self.CHART[1],
            before="Arial", after="Calibri",
        ))
        with pytest.raises(ApplyError, match="native chart"):
            apply_changes(adversarial_deck, cs, tmp_path / "nope.pptx")


class TestTheCostOfAnApplyDoesNotGrowWithTheChangeCount:
    """The guards used to run per change, and each one scanned the package.

    Every slide, every rels file, every chart part — the same scan, once for
    every change in the set. Profiled on a 272-change tidy of a 26-slide deck:
    **32,778 reads of the archive, and 102 of the 103 seconds it took**. A
    1,199-change tidy on a 41-slide deck took 7 minutes 48 seconds, and the cost
    *per change* was rising, because the scan is the same size however many
    changes there are.

    Hoisted out of the loop: 103s → 1.5s on the first, 468s → 3.9s on the
    second. The whole test suite halved.

    The number is not what this test protects — machines differ. It protects the
    shape: the work is done once, so putting the scan back inside the loop fails
    here rather than being noticed a year later on somebody's 300-slide deck.
    """

    def _tidy_changes(self, deck, count):
        info = inspect(deck)
        shape = next(s for sl in info.slides for s in sl.shapes
                     if s.kind not in ("table", "chart") and s.runs
                     and any(r.text.strip() for r in s.runs))
        slide = next(sl.number for sl in info.slides for s in sl.shapes if s is shape)
        run = next(r for r in shape.runs if r.text.strip())
        cs = ChangeSet(deck=str(deck))
        # The same edit repeated: only the first can apply, and the rest fail
        # honestly. What is being counted is the guard work, which happens
        # before any of them are written.
        for i in range(count):
            cs.add(Change(id=f"c{i}", op=Op.SET_TEXT, slide=slide,
                          target=shape.id, before=run.text, after=run.text + "."))
        cs.approve_all()
        return cs

    def _reads_during(self, deck, changeset, out, monkeypatch):
        from slide_wright.package import Package

        calls = []
        original = Package.read
        monkeypatch.setattr(
            Package, "read",
            lambda self, name: (calls.append(name), original(self, name))[1],
        )
        try:
            apply_changes(deck, changeset, out)
        except Exception:
            pass  # a refusal still did the guard work, which is what is counted
        return len(calls)

    def test_ten_times_the_changes_is_not_ten_times_the_work(
        self, adversarial_deck, tmp_path, monkeypatch
    ):
        few = self._reads_during(
            adversarial_deck, self._tidy_changes(adversarial_deck, 2),
            tmp_path / "few.pptx", monkeypatch)
        many = self._reads_during(
            adversarial_deck, self._tidy_changes(adversarial_deck, 20),
            tmp_path / "many.pptx", monkeypatch)

        assert many < few * 2, (
            f"{few} reads for 2 changes and {many} for 20 — the package is "
            "being rescanned per change again"
        )

    def test_the_refusal_still_happens_and_still_explains_itself(
        self, adversarial_deck, tmp_path
    ):
        """The lookup is cheap; the message is not. It must still be produced
        for the one change that is actually refused."""
        deck = inspect(adversarial_deck)
        chart = next((s.number, x) for s in deck.slides for x in s.shapes
                     if x.kind == "chart")
        cs = approved(adversarial_deck, Change(
            id="c1", op=Op.SET_FONT, slide=chart[0], target=chart[1].id,
            before="Arial", after="Calibri",
        ))
        with pytest.raises(ApplyError, match="native chart"):
            apply_changes(adversarial_deck, cs, tmp_path / "no.pptx")


class TestTwoWritesAtOnceDoNotCollide:
    """The scratch file has to be unique per call, not per process.

    Keying it on the process id was enough for two processes and not for two
    threads, and two threads is the ordinary case — a server handling two
    requests. Three concurrent applies produced `PermissionError: the process
    cannot access the file because it is being used by another process`: one
    success and two raw 500s, with the work in them lost.
    """

    def test_concurrent_writes_to_different_outputs_all_succeed(
        self, adversarial_deck, tmp_path
    ):
        import threading

        from slide_wright.apply import _write_package

        results: list[tuple[int, str]] = []

        def write(i: int) -> None:
            try:
                _write_package(
                    adversarial_deck, tmp_path / f"out-{i}.pptx",
                    {"ppt/slides/slide1.xml": f"<p:sld n='{i}'/>".encode()},
                )
                results.append((i, "ok"))
            except Exception as exc:  # noqa: BLE001
                results.append((i, f"{type(exc).__name__}: {exc}"))

        threads = [threading.Thread(target=write, args=(i,)) for i in range(6)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(30)

        assert [r for _, r in results] == ["ok"] * 6, results

    def test_each_output_is_a_whole_package(self, adversarial_deck, tmp_path):
        import threading

        from slide_wright.apply import _write_package

        def write(i: int) -> None:
            _write_package(adversarial_deck, tmp_path / f"out-{i}.pptx", {})

        threads = [threading.Thread(target=write, args=(i,)) for i in range(6)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(30)

        expected = Package.open(adversarial_deck).part_count
        for i in range(6):
            assert Package.open(tmp_path / f"out-{i}.pptx").part_count == expected

    def test_nothing_is_left_behind(self, adversarial_deck, tmp_path):
        from slide_wright.apply import _write_package

        _write_package(adversarial_deck, tmp_path / "out.pptx", {})
        assert not list(tmp_path.glob("*.partial"))
        assert not list(tmp_path.glob(".*"))


class TestAnEditKeepsTheFormattingItDidNotAskAbout:
    """The promise, at run granularity.

    "Change what I asked, preserve everything else" is checkable inside a single
    paragraph, and that is where it was failing. Replacing two words at the
    front of a sentence used to write the whole joined string into the first run
    and blank the rest, so the bold on a figure three words later disappeared —
    applied, verified, and invisible to a content diff because the text was
    identical.
    """

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def _paragraph(self, *runs: tuple[str, str]):
        from lxml import etree

        body = "".join(
            f'<a:r><a:rPr {attrs}/><a:t>{text}</a:t></a:r>' for text, attrs in runs
        )
        return etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f"<a:p>{body}</a:p></p:sp>"
        )

    def _runs(self, shape):
        return [
            (r.find(f"{self.A}t").text or "", r.find(f"{self.A}rPr").get("b"))
            for r in shape.findall(f".//{self.A}r")
        ]

    def test_a_bold_figure_survives_an_edit_to_the_words_before_it(self):
        from slide_wright.apply import _set_text

        shape = self._paragraph(
            ("Revenue grew ", 'b="0"'), ("15%", 'b="1"'), (" in FY25", 'b="0"')
        )
        assert _set_text(shape, "Revenue grew", "Revenue increased")
        assert self._runs(shape) == [
            ("Revenue increased ", "0"),
            ("15%", "1"),
            (" in FY25", "0"),
        ], "the bold run was neither emptied nor merged into its neighbour"

    def test_a_run_after_the_span_keeps_its_own_text(self):
        from slide_wright.apply import _set_text

        shape = self._paragraph(("As at ", 'b="0"'), ("Q3", 'b="1"'), (" close", 'b="0"'))
        assert _set_text(shape, "As at Q3", "As at Q4")
        assert self._runs(shape) == [("As at Q4", "0"), ("", "1"), (" close", "0")], (
            "only the runs the span covers may change"
        )

    def test_a_span_ending_mid_run_keeps_the_remainder(self):
        from slide_wright.apply import _set_text

        shape = self._paragraph(("Fees ", 'b="0"'), ("rose sharply", 'b="1"'))
        assert _set_text(shape, "Fees rose", "Fees fell")
        assert self._runs(shape) == [("Fees fell", "0"), (" sharply", "1")], (
            "the tail beyond the span keeps its own formatting"
        )

    def test_an_edit_spanning_two_bullets_leaves_the_third_alone(self, tmp_path):
        """The visible version: text collapsing out of the paragraphs it lived in."""
        from lxml import etree

        from slide_wright.apply import _set_text

        paras = "".join(
            f'<a:p><a:r><a:rPr/><a:t>{t}</a:t></a:r></a:p>'
            for t in ("Margin held at 42%", "Headcount fell to 310", "Cash runway 18 months")
        )
        shape = etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            f' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{paras}</p:sp>'
        )
        assert _set_text(shape, "42%\nHeadcount fell", "42%\nHeadcount rose")
        third = shape.findall(f"{self.A}p")[2]
        assert third.find(f".//{self.A}t").text == "Cash runway 18 months", (
            "a bullet the span never reached must still hold its own text"
        )

    def test_it_holds_through_save_and_reload(self, tmp_path):
        """End to end on a real package, because XML in memory is not a deck.

        Built here rather than taken from the corpus: the adversarial deck has
        no shape carrying two differently formatted runs in one paragraph, which
        is exactly the arrangement this defect needed. That absence is why every
        existing narrowness test passed while the promise was broken.
        """
        from pptx import Presentation
        from pptx.util import Inches, Pt

        src = tmp_path / "mixed.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text, bold in (("Revenue grew ", False), ("15%", True), (" in FY25", False)):
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(18)
        prs.save(str(src))

        shape = next(s for s in inspect(src).slides[0].shapes if s.text.strip())
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
            before="Revenue grew", after="Revenue increased",
        ))
        result = apply_changes(src, cs, out)
        assert not result.failed, result.failed

        after = next(s for s in inspect(out).slides[0].shapes if s.text.strip())
        assert [(r.text, r.bold) for r in after.runs] == [
            ("Revenue increased ", False),
            ("15%", True),
            (" in FY25", False),
        ], "the emphasis on the figure must survive an edit that never named it"


class TestACellSplitAcrossRunsIsStillOneValue:
    """Where the figures live, and where a wrong one arrives with a citation.

    A cell's value is one thing to a reader and often several runs to OOXML:
    part of a number gets bolded, or a language boundary falls inside it, and
    "1,234" is stored as "1,2" + "34". Replacing it used to write the new value
    into the first run and leave the rest, so a refresh cited to a spreadsheet
    coordinate produced 1,98734 and reported it applied.
    """

    def _deck_with_a_split_cell(self, tmp_path, pieces=(("1,2", False), ("34", True))):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        path = tmp_path / "split-cell.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(
            2, 2, Inches(1), Inches(1), Inches(6), Inches(1.5)
        ).table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "FY25"
        table.cell(1, 0).text = "Revenue"
        paragraph = table.cell(1, 1).text_frame.paragraphs[0]
        paragraph.text = ""
        for text, bold in pieces:
            run = paragraph.add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(14)
        prs.save(str(path))
        return path

    def _table(self, deck):
        return next(s for s in inspect(deck).slides[0].shapes if s.kind == "table")

    def test_replacing_a_split_figure_leaves_one_figure(self, tmp_path):
        src = self._deck_with_a_split_cell(tmp_path)
        table = self._table(src)
        assert table.cell(1, 1) == "1,234"

        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=1, target=f"{table.id}/r1/c1",
            before="1,234", after="1,987", citation="q3.xlsx!B2",
        ))
        result = apply_changes(src, cs, out)
        assert not result.failed, result.failed
        assert self._table(out).cell(1, 1) == "1,987", (
            "the old value's tail must not survive alongside the new one"
        )

    def test_a_wrong_before_is_refused_rather_than_written_anywhere(self, tmp_path):
        """The safety half: a stale `before` must not become a blind write."""
        src = self._deck_with_a_split_cell(tmp_path)
        table = self._table(src)
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=1, target=f"{table.id}/r1/c1",
            before="9,999", after="1,987",
        ))
        result = apply_changes(src, cs, out)
        assert len(result.failed) == 1
        assert "9,999" in str(result.failed[0][1]), "the refusal must name the value it looked for"
        assert self._table(out).cell(1, 1) == "1,234", "the cell must be untouched"

    def test_a_cell_with_no_before_is_set_outright(self, tmp_path):
        """`before=None` means "this cell now says X" — the whole cell is the target."""
        src = self._deck_with_a_split_cell(tmp_path)
        table = self._table(src)
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_TABLE_CELL, slide=1, target=f"{table.id}/r1/c1",
            before=None, after="n/a",
        ))
        assert not apply_changes(src, cs, out).failed
        assert self._table(out).cell(1, 1) == "n/a"


class TestATargetNamingOneRunChangesOneRun:
    """Three formatting ops, one addressing rule.

    `<shape>/run/<index>` is how conformance work stays checkable: the word that
    was pasted in the wrong font comes back fixed and every other run byte
    identical. Font and colour honoured it; size did not, and set `sz` on every
    run in the shape whatever the target said.
    """

    def _deck(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        path = tmp_path / "sizes.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text, size in (("Headline ", 28), ("footnote", 10), (" tail", 28)):
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(size)
        prs.save(str(path))
        return path

    def _runs(self, deck):
        shape = next(s for s in inspect(deck).slides[0].shapes if s.text.strip())
        return shape.id, [(r.text, r.size_pt) for r in shape.runs]

    def test_resizing_one_run_leaves_the_headline_alone(self, tmp_path):
        src = self._deck(tmp_path)
        shape_id, _ = self._runs(src)
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_FONT_SIZE, slide=1, target=f"{shape_id}/run/1",
            before=10, after=12,
        ))
        assert not apply_changes(src, cs, out).failed
        assert self._runs(out)[1] == [
            ("Headline ", 28.0), ("footnote", 12.0), (" tail", 28.0)
        ], "a change addressed at one run must not resize its neighbours"

    def test_a_bare_shape_id_still_resizes_the_whole_shape(self, tmp_path):
        src = self._deck(tmp_path)
        shape_id, _ = self._runs(src)
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_FONT_SIZE, slide=1, target=shape_id, after=12,
        ))
        assert not apply_changes(src, cs, out).failed
        assert {size for _, size in self._runs(out)[1]} == {12.0}

    def test_a_run_index_off_the_end_is_refused(self, tmp_path):
        src = self._deck(tmp_path)
        shape_id, before = self._runs(src)
        out = tmp_path / "o.pptx"
        cs = approved(src, Change(
            id="c1", op=Op.SET_FONT_SIZE, slide=1, target=f"{shape_id}/run/9",
            after=12,
        ))
        assert len(apply_changes(src, cs, out).failed) == 1
        assert self._runs(out)[1] == before, "a refused change must write nothing"


class TestNarrownessAskedOfEveryShape:
    """The property, swept rather than exemplified.

    Adding a split-run slide to the corpus caught nothing on its own: 884 tests
    passed over it with the applier still broken, because every one of them
    asked about a shape it had chosen. A construct nothing interrogates is a
    construct nothing checks, which is the same lesson the surfaces taught and
    it applies to fixtures too.

    These two sweep the invariant over whatever the deck happens to contain, so
    a fixture added later is covered by the question the day it lands.
    """

    #: Every attribute `TextRun` records. The list used to be six of them, and
    #: was written when six was all there were -- so as underline, strikethrough,
    #: baseline, capitals and the hyperlink arrived, the sweep that asks "did
    #: anything else about this run change" quietly stopped asking about half of
    #: what a run is. Read off the dataclass rather than listed, so the next
    #: attribute is covered the day it lands instead of the day someone
    #: remembers this line.
    FIELDS = tuple(f.name for f in fields(TextRun) if f.name != "paragraph")

    def _fingerprint(self, shape):
        return [tuple(getattr(r, f) for f in self.FIELDS) for r in shape.runs]

    def test_editing_one_run_leaves_every_other_run_untouched(
        self, adversarial_deck, tmp_path
    ):
        deck = inspect(adversarial_deck)
        swept = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.kind in {"table", "group"} or len(shape.runs) < 2:
                    continue
                head = shape.runs[0].text
                if len(head) < 2:
                    continue
                # A proper prefix, so no run matches exactly and the edit goes
                # through the joining path rather than the single-run one.
                before = head[: len(head) - 1]
                out = tmp_path / f"s{slide.number}-{shape.id}.pptx"
                cs = approved(adversarial_deck, Change(
                    id="c1", op=Op.SET_TEXT, slide=slide.number, target=shape.id,
                    before=before, after=before.upper(),
                ))
                result = apply_changes(adversarial_deck, cs, out)
                assert not result.failed, result.failed
                after = next(
                    s for s in inspect(out).slides[slide.number - 1].shapes
                    if s.id == shape.id
                )
                assert self._fingerprint(after)[1:] == self._fingerprint(shape)[1:], (
                    f"slide {slide.number} shape {shape.id}: an edit inside run 1 "
                    f"changed a run it never addressed"
                )
                swept += 1
        assert swept, "the corpus must contain a shape with more than one run"

    def test_a_replaced_cell_holds_exactly_the_new_value(
        self, adversarial_deck, tmp_path
    ):
        deck = inspect(adversarial_deck)
        swept = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.kind != "table":
                    continue
                for address, value in shape.table_cells.items():
                    if not value:
                        continue
                    out = tmp_path / f"c{slide.number}-{shape.id}-{address.replace('/', '')}.pptx"
                    cs = approved(adversarial_deck, Change(
                        id="c1", op=Op.SET_TABLE_CELL, slide=slide.number,
                        target=f"{shape.id}/{address}", before=value, after="ZZ",
                    ))
                    result = apply_changes(adversarial_deck, cs, out)
                    assert not result.failed, result.failed
                    written = next(
                        s for s in inspect(out).slides[slide.number - 1].shapes
                        if s.id == shape.id
                    ).table_cells[address]
                    assert written == "ZZ", (
                        f"slide {slide.number} table {shape.id} cell {address}: "
                        f"replacing {value!r} left {written!r}"
                    )
                    swept += 1
        assert swept, "the corpus must contain a table with a non-empty cell"


class TestARunTargetMeansWhatItMeantWhenApproved:
    """Indices are read against the deck the reviewer saw, not the one mid-edit.

    A run index names a position in `inspect`'s enumeration, which skips runs
    with no text. A text edit can empty a run, so a change set containing both a
    text edit and a run-addressed formatting change on the same shape shifts its
    own indices while it is being applied.
    """

    def _five_runs(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        path = tmp_path / "five.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text in ("AAA", "BBB", "CCC", "DDD", "EEE"):
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(18)
        prs.save(str(path))
        return path

    def _fonts(self, deck):
        shape = next(s for s in inspect(deck).slides[0].shapes if s.text.strip())
        return [(r.text, r.font) for r in shape.runs]

    def test_an_earlier_text_edit_does_not_move_a_later_run_target(self, tmp_path):
        src = self._five_runs(tmp_path)
        shape_id = next(s for s in inspect(src).slides[0].shapes if s.text.strip()).id
        out = tmp_path / "o.pptx"
        cs = approved(
            src,
            # Empties the second run, so the enumeration shrinks by one.
            Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape_id,
                   before="AAABBB", after="X"),
            # Run 3 is DDD in the deck as read; it was EEE by the time this landed.
            Change(id="c2", op=Op.SET_FONT, slide=1, target=f"{shape_id}/run/3",
                   before="Arial", after="Georgia"),
        )
        result = apply_changes(src, cs, out)
        assert not result.failed, result.failed
        assert self._fonts(out) == [
            ("X", "Arial"), ("CCC", "Arial"), ("DDD", "Georgia"), ("EEE", "Arial")
        ], "the run the change named must be the run that changed"

    def test_the_same_holds_for_size(self, tmp_path):
        src = self._five_runs(tmp_path)
        shape_id = next(s for s in inspect(src).slides[0].shapes if s.text.strip()).id
        out = tmp_path / "o.pptx"
        cs = approved(
            src,
            Change(id="c1", op=Op.SET_TEXT, slide=1, target=shape_id,
                   before="AAABBB", after="X"),
            Change(id="c2", op=Op.SET_FONT_SIZE, slide=1,
                   target=f"{shape_id}/run/3", before=18, after=9),
        )
        assert not apply_changes(src, cs, out).failed
        shape = next(s for s in inspect(out).slides[0].shapes if s.text.strip())
        assert [(r.text, r.size_pt) for r in shape.runs] == [
            ("X", 18.0), ("CCC", 18.0), ("DDD", 9.0), ("EEE", 18.0)
        ]


class TestAFormattingLockSurvivesATextEdit:
    """The guarantee, rather than the mechanism that happened to break it.

    `formatting` means *"fix the words, leave my styling exactly as it is"*. It
    is expressed as a refusal of SET_FONT, SET_COLOR and SET_FONT_SIZE, which is
    correct about what the user can ask for and says nothing about what the
    applier does on its way through. A text edit is permitted under this lock by
    design -- fixing words is the whole point of holding it -- and the applier
    used to strip the styling of every run it joined while doing so.

    So the lock could be held, honoured by the change set, reported as honoured,
    and the styling gone anyway. A guarantee is only worth what the writer does,
    not what the gate refuses.
    """

    def test_the_styling_is_still_there_afterwards(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        from slide_wright.changeset import Lock

        src = tmp_path / "locked.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        for text, bold, size in (
            ("Our margin held at ", False, 18),
            ("42%", True, 24),
            (" through the quarter", False, 18),
        ):
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(size)
        prs.save(str(src))

        shape = next(s for s in inspect(src).slides[0].shapes if s.text.strip())
        styling = [(r.bold, r.size_pt) for r in shape.runs]

        cs = ChangeSet(deck=str(src))
        cs.lock("formatting", reason="leave my styling exactly as it is")
        cs.add(Change(
            id="c1", op=Op.SET_TEXT, slide=1, target=shape.id,
            before="Our margin held at", after="Our margin improved to",
        ))
        cs.approve_all()

        out = tmp_path / "o.pptx"
        result = apply_changes(src, cs, out)
        assert not result.failed, result.failed

        after = next(s for s in inspect(out).slides[0].shapes if s.text.strip())
        assert [(r.bold, r.size_pt) for r in after.runs] == styling, (
            "a formatting lock has to survive the edit it was held across"
        )
        assert "42%" in after.text


class TestEveryLockIsAPropertyOfTheWrittenDeck:
    """Nine scopes, each stated as something true of the output file.

    A lock is tested elsewhere by what the change set does with it: the change
    arrives, the lock rejects it, the status reads REJECTED. That is the gate
    working, and it is not the guarantee. The guarantee is about the deck that
    comes out, and the difference is exactly where this product keeps finding
    defects -- a `formatting` lock permits text edits by design, and the applier
    used to strip formatting while making one.

    So each case holds a lock, sends one change the lock must refuse *and one it
    must allow*, applies, and asks the output whether the protected property
    survived the edit it was held across. A lock with nothing applied alongside
    it proves nothing: the deck is never rewritten, so of course it did not
    change.
    """

    def _find(self, deck, kind):
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.kind == kind:
                    return slide.number, shape
        raise AssertionError(f"the corpus has no {kind}")

    def _text_shape(self, deck, slide_number):
        return next(
            s
            for s in deck.slides[slide_number - 1].shapes
            if s.kind not in {"table", "group", "chart", "picture"} and s.text.strip()
        )

    # -- the properties, read off a deck --------------------------------------

    def _all_text(self, deck):
        return {(sl.number, s.id): s.text for sl in deck.slides for s in sl.shapes}

    def _all_geometry(self, deck):
        return {
            (sl.number, s.id): (s.x, s.y, s.cx, s.cy)
            for sl in deck.slides
            for s in sl.shapes
        }

    def _all_formatting(self, deck):
        return {
            (sl.number, s.id): [
                (r.text, r.size_pt, r.bold, r.italic, r.font, r.color) for r in s.runs
            ]
            for sl in deck.slides
            for s in sl.shapes
        }

    def _all_styling(self, deck):
        """Formatting with the text removed, so a text edit can be applied to the
        same shape the property is asserted over. Comparing text too would make
        the permitted change look like a violation and force the shape to be
        excluded -- which is what made the first version of this suite pass
        against an applier that destroyed styling."""
        return {
            (sl.number, s.id): [
                (r.size_pt, r.bold, r.italic, r.font, r.color) for r in s.runs
            ]
            for sl in deck.slides
            for s in sl.shapes
        }

    def _all_cells(self, deck):
        return {
            (sl.number, s.id): dict(s.table_cells)
            for sl in deck.slides
            for s in sl.shapes
            if s.kind == "table"
        }

    def _all_digits(self, deck):
        text = " ".join(s.text for sl in deck.slides for s in sl.shapes)
        return sorted(re.findall(r"\d+", text))

    # -- the harness ----------------------------------------------------------

    def _apply_under(self, deck_path, tmp_path, scope, target, forbidden, permitted):
        cs = ChangeSet(deck=str(deck_path))
        cs.lock(scope, target)
        cs.add(forbidden)
        cs.add(permitted)
        cs.approve_all()
        assert forbidden.status is Status.REJECTED, (
            f"the {scope} lock did not refuse the change it exists to refuse"
        )
        out = tmp_path / f"{scope}.pptx"
        result = apply_changes(deck_path, cs, out)
        assert not result.failed, result.failed
        assert result.applied == [permitted], "the permitted change must land"
        return inspect(out)

    def _multi_run_shape(self, deck, slide_number):
        """The shape whose text is several runs — the one narrowness needs."""
        return next(
            s
            for s in deck.slides[slide_number - 1].shapes
            if s.kind not in {"table", "group", "chart", "picture"} and len(s.runs) > 1
        )

    def _wordy(self, deck):
        """A digit-free text edit that every case can use as its permitted change."""
        shape = self._text_shape(deck, 1)
        return shape, Change(
            id="p", op=Op.SET_TEXT, slide=1, target=shape.id,
            before=shape.runs[0].text, after="Rewritten opening",
        )

    # -- the nine -------------------------------------------------------------

    def test_a_slide_lock_leaves_that_slide_byte_identical(
        self, adversarial_deck, tmp_path
    ):
        deck = inspect(adversarial_deck)
        protected = self._text_shape(deck, 1)
        elsewhere = self._text_shape(deck, 6)
        cs = ChangeSet(deck=str(adversarial_deck))
        cs.lock("slide", "1")
        forbidden = cs.add(Change(
            id="f", op=Op.SET_TEXT, slide=1, target=protected.id,
            before=protected.runs[0].text, after="CHANGED",
        ))
        permitted = cs.add(Change(
            id="p", op=Op.SET_TEXT, slide=6, target=elsewhere.id,
            before=elsewhere.runs[0].text, after="edited",
        ))
        cs.approve_all()
        assert forbidden.status is Status.REJECTED

        out = tmp_path / "slide.pptx"
        assert not apply_changes(adversarial_deck, cs, out).failed
        assert permitted.status is Status.APPLIED

        part = deck.slides[0].part_name
        assert zipfile.ZipFile(out).read(part) == zipfile.ZipFile(
            adversarial_deck
        ).read(part), "a locked slide's part must come back byte for byte"

    def test_a_shape_lock_leaves_that_shape_alone(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        protected = self._multi_run_shape(deck, 7)
        # On the *same slide*, so the part is rewritten around the lock rather
        # than left alone. A locked shape on a part nothing touched is not a
        # test of the lock.
        neighbour = next(
            s for s in deck.slides[6].shapes
            if s.id != protected.id and s.kind == "shape" and s.text.strip()
        )
        after = self._apply_under(
            adversarial_deck, tmp_path, "shape", protected.id,
            Change(
                id="f", op=Op.SET_TEXT, slide=7,
                target=f"{protected.id}/run/0",
                before=protected.runs[0].text, after="CHANGED",
            ),
            Change(
                id="p", op=Op.SET_TEXT, slide=7, target=neighbour.id,
                before=neighbour.runs[0].text[:6], after="EDITED",
            ),
        )
        survived = next(s for s in after.slides[6].shapes if s.id == protected.id)
        assert survived.text == protected.text
        assert [r.bold for r in survived.runs] == [r.bold for r in protected.runs]

    def test_a_numbers_lock_leaves_every_figure_where_it_was(
        self, adversarial_deck, tmp_path
    ):
        deck = inspect(adversarial_deck)
        slide_number, table = self._find(deck, "table")
        shape, permitted = self._wordy(deck)
        assert not re.search(r"\d", shape.runs[0].text), (
            "the permitted edit has to be digit-free or the lock would refuse it too"
        )
        after = self._apply_under(
            adversarial_deck, tmp_path, "numbers", "",
            Change(
                id="f", op=Op.SET_TABLE_CELL, slide=slide_number,
                target=f"{table.id}/r1/c1", before=table.cell(1, 1), after="99.9x",
            ),
            permitted,
        )
        assert self._all_digits(after) == self._all_digits(deck)

    def test_a_wording_lock_leaves_every_word(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        slide_number, table = self._find(deck, "table")
        shape = self._text_shape(deck, 1)
        after = self._apply_under(
            adversarial_deck, tmp_path, "wording", "",
            Change(
                id="f", op=Op.SET_TEXT, slide=1, target=shape.id,
                before=shape.runs[0].text, after="CHANGED",
            ),
            Change(
                id="p", op=Op.MOVE, slide=slide_number, target=table.id,
                before=[table.x, table.y], after=[table.x + 100000, table.y],
            ),
        )
        assert self._all_text(after) == self._all_text(deck)

    def test_a_layout_lock_leaves_every_box(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        slide_number, table = self._find(deck, "table")
        _, permitted = self._wordy(deck)
        after = self._apply_under(
            adversarial_deck, tmp_path, "layout", "",
            Change(
                id="f", op=Op.MOVE, slide=slide_number, target=table.id,
                before=[table.x, table.y], after=[table.x + 100000, table.y],
            ),
            permitted,
        )
        assert self._all_geometry(after) == self._all_geometry(deck)

    def test_a_formatting_lock_survives_the_text_edit_it_permits(
        self, adversarial_deck, tmp_path
    ):
        deck = inspect(adversarial_deck)
        styled = self._multi_run_shape(deck, 7)
        head = styled.runs[0].text
        after = self._apply_under(
            adversarial_deck, tmp_path, "formatting", "",
            Change(
                id="f", op=Op.SET_FONT, slide=7, target=f"{styled.id}/run/1",
                before=styled.runs[1].font, after="Georgia",
            ),
            # The permitted edit lands on the styled shape, through the joining
            # path -- a prefix, so no run matches exactly. This is the whole
            # point: "fix the words" is what the lock is held across.
            Change(
                id="p", op=Op.SET_TEXT, slide=7, target=styled.id,
                before=head[: len(head) - 1], after=head[: len(head) - 1].upper(),
            ),
        )
        assert self._all_styling(after) == self._all_styling(deck), (
            "a formatting lock has to survive the text edit it permits"
        )

    def test_a_tables_lock_leaves_every_cell(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        slide_number, table = self._find(deck, "table")
        _, permitted = self._wordy(deck)
        after = self._apply_under(
            adversarial_deck, tmp_path, "tables", "",
            Change(
                id="f", op=Op.SET_TABLE_CELL, slide=slide_number,
                target=f"{table.id}/r1/c1", before=table.cell(1, 1), after="ZZ",
            ),
            permitted,
        )
        assert self._all_cells(after) == self._all_cells(deck)

    def test_a_charts_lock_leaves_the_chart_frame(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        slide_number, chart = self._find(deck, "chart")
        _, permitted = self._wordy(deck)
        after = self._apply_under(
            adversarial_deck, tmp_path, "charts", "",
            Change(
                id="f", op=Op.MOVE, slide=slide_number, target=chart.id,
                before=[chart.x, chart.y], after=[chart.x + 100000, chart.y],
                object_kind="chart",
            ),
            permitted,
        )
        held = next(s for s in after.slides[slide_number - 1].shapes if s.id == chart.id)
        assert (held.x, held.y) == (chart.x, chart.y)

    def test_a_media_lock_leaves_the_picture(self, adversarial_deck, tmp_path):
        deck = inspect(adversarial_deck)
        slide_number, picture = self._find(deck, "picture")
        _, permitted = self._wordy(deck)
        after = self._apply_under(
            adversarial_deck, tmp_path, "media", "",
            Change(
                id="f", op=Op.MOVE, slide=slide_number, target=picture.id,
                before=[picture.x, picture.y],
                after=[picture.x + 100000, picture.y], object_kind="picture",
            ),
            permitted,
        )
        held = next(
            s for s in after.slides[slide_number - 1].shapes if s.id == picture.id
        )
        assert (held.x, held.y) == (picture.x, picture.y)


class TestAShapeEndsWithTheLinesItWasAskedFor:
    """The line a reader counts, not the line a run walk can see.

    A replacement crossing paragraphs used to write the whole new string into
    the first run and empty every other one. That is right about the words and
    wrong about the shape: the emptied paragraphs stay, each still drawing the
    bullet it inherits from the layout, so a four-line placeholder set to one
    line came back as one line and three blank bullets.

    Nothing could see it. `runs` holds only runs with text, so `ShapeInfo.text`
    read back clean and the diff said *0 change how it looks*. And it is the
    ordinary path, not a corner: `SetSpec` carries no `before`, so every typed
    edit in the workspace replaces the whole shape.
    """

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def _shape(self, *lines: str, bullet: bool = False):
        from lxml import etree

        pPr = '<a:pPr><a:buChar char="\u2022"/></a:pPr>' if bullet else ""
        paras = "".join(
            f"<a:p>{pPr}<a:r><a:rPr/><a:t>{line}</a:t></a:r></a:p>" for line in lines
        )
        return etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            f' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{paras}</p:sp>'
        )

    def _lines(self, shape) -> list[str]:
        return [
            "".join(t.text or "" for t in para.findall(f".//{self.A}t"))
            for para in shape.findall(f"{self.A}p")
        ]

    def test_four_lines_set_to_one_leaves_one(self):
        from slide_wright.apply import _set_text

        shape = self._shape("Alpha", "Beta", "Gamma", "Delta")
        assert _set_text(shape, "Alpha\nBeta\nGamma\nDelta", "One line")
        assert self._lines(shape) == ["One line"], (
            "the paragraphs the replacement no longer needs must go, not be emptied"
        )

    def test_a_newline_in_the_replacement_becomes_a_paragraph(self):
        """The same mistake facing the other way.

        A newline character inside `<a:t>` is not a line break in OOXML, and
        reading it back gives a string the engine cannot tell from two
        paragraphs -- so the engine could not distinguish its own two outcomes.
        """
        from slide_wright.apply import _set_text

        shape = self._shape("Alpha", "Beta", "Gamma")
        assert _set_text(shape, "Alpha\nBeta\nGamma", "First\nSecond")
        assert self._lines(shape) == ["First", "Second"]
        assert not any(
            "\n" in (t.text or "") for t in shape.findall(f".//{self.A}t")
        ), "a line break must be a paragraph, never a character in a run"

    def test_a_line_added_keeps_the_bullet_of_the_line_above(self):
        from slide_wright.apply import _set_text

        shape = self._shape("Alpha", "Beta", bullet=True)
        assert _set_text(shape, "Alpha\nBeta", "Alpha\nBeta\nGamma")
        assert self._lines(shape) == ["Alpha", "Beta", "Gamma"]
        assert all(
            para.find(f"{self.A}pPr/{self.A}buChar") is not None
            for para in shape.findall(f"{self.A}p")
        ), "a line added to a bulleted list means a bullet"

    def test_a_line_added_after_a_linked_one_points_nowhere(self):
        """Mine, from this fix: an edit must not invent a hyperlink either.

        The new paragraph is cloned from the line above so the bullet and the
        styling carry. A link is not styling -- it is what the deck *does* --
        and cloning one made a line the user typed point at a target they never
        named. The same category of wrong as an edit destroying a link, from the
        other direction.
        """
        from lxml import etree

        from slide_wright.apply import _set_text

        shape = etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            "<a:p><a:r><a:rPr/><a:t>Alpha</a:t></a:r></a:p>"
            '<a:p><a:r><a:rPr><a:hlinkClick r:id="rId9"/></a:rPr>'
            "<a:t>Beta</a:t></a:r></a:p></p:sp>"
        )
        assert _set_text(shape, "Alpha\nBeta", "Alpha\nBeta\nGamma")
        paragraphs = shape.findall(f"{self.A}p")
        assert [p.find(f".//{self.A}t").text for p in paragraphs] == [
            "Alpha", "Beta", "Gamma",
        ]
        assert paragraphs[1].find(f".//{self.A}hlinkClick") is not None, (
            "the line that had the link keeps it"
        )
        assert paragraphs[2].find(f".//{self.A}hlinkClick") is None, (
            "the line that was added must not have acquired one"
        )

    def test_a_line_added_after_a_footer_does_not_copy_the_slide_number(self):
        """Also mine: a field is a second thing a clone must not carry.

        A slide number, a date, a footer -- 156 paragraphs across 6 of the 26
        real decks hold one. Cloning a line that has one puts a second copy of
        it on the slide, and an explicit line break comes with it. Both are
        things the deck *does* rather than how it looks.
        """
        from lxml import etree

        from slide_wright.apply import _set_text

        shape = etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            "<a:p><a:r><a:rPr/><a:t>Alpha</a:t></a:r></a:p>"
            "<a:p><a:r><a:rPr/><a:t>Beta</a:t></a:r><a:br/>"
            '<a:fld id="{1}" type="slidenum"><a:t>7</a:t></a:fld></a:p></p:sp>'
        )
        assert _set_text(shape, "Alpha\nBeta", "Alpha\nBeta\nGamma")
        paragraphs = shape.findall(f"{self.A}p")
        assert paragraphs[1].find(f"{self.A}fld") is not None, "the footer keeps its field"
        assert paragraphs[2].find(f"{self.A}fld") is None, (
            "the line that was added must not carry a second slide number"
        )
        assert paragraphs[2].find(f"{self.A}br") is None, (
            "nor a line break the user never typed"
        )

    def test_a_line_the_span_never_reached_keeps_its_own_text(self):
        from slide_wright.apply import _set_text

        shape = self._shape("Alpha", "Beta", "Gamma", "Delta")
        assert _set_text(shape, "Alpha\nBeta\nGamma", "One")
        assert self._lines(shape) == ["One", "Delta"]

    def test_the_tail_of_the_last_line_survives_the_paragraphs_below_it(self):
        from slide_wright.apply import _set_text

        shape = self._shape("one two", "three four")
        assert _set_text(shape, "two\nthree", "TWO")
        assert self._lines(shape) == ["one TWO four"]

    def test_a_tail_travels_with_the_last_line_when_lines_are_added(self):
        from slide_wright.apply import _set_text

        shape = self._shape("one two", "three four")
        assert _set_text(shape, "two\nthree", "TWO\nMID\nTHREE")
        assert self._lines(shape) == ["one TWO", "MID", "THREE four"]

    def test_a_replacement_starting_on_a_boundary_is_not_dropped(self):
        """`lo == hi` is an insertion point, not a span of nothing.

        Mine, from the first version of this fix: a replacement covering no
        character of the paragraph it starts in wrote nothing and then removed
        the paragraph below, so the deck lost a line and gained no replacement.
        """
        from slide_wright.apply import _set_text

        shape = self._shape("abc", "def")
        assert _set_text(shape, "\ndef", "X")
        assert self._lines(shape) == ["abcX"]

    def test_the_only_paragraph_is_emptied_rather_than_removed(self):
        """A text body must hold at least one paragraph, or the deck is broken."""
        from slide_wright.apply import _replace_across_paragraphs
        from lxml import etree

        shape = self._shape("Alpha")
        paragraphs = shape.findall(f"{self.A}p")
        assert _replace_across_paragraphs(paragraphs, "Alpha", "")
        assert len(shape.findall(f"{self.A}p")) == 1
        assert self._lines(shape) == [""]


class TestBlankLinesAreVisibleEndToEnd:
    """On a real package, because XML in memory is not a deck.

    The shape this needs -- a body placeholder with bullets in it -- was not in
    the corpus until this defect was found. It had a table, whose cells are
    paragraphs, and a group, whose children hold one each. Neither is a text
    body with lines in it, so the question had never been put to a deck.
    """

    def _bulleted(self, deck):
        """The first ordinary text shape holding more than one line.

        A group is excluded on purpose: its lines live in separate text bodies,
        one per child, and what happens to those is a different guarantee --
        recorded below.
        """
        info = inspect(deck)
        return next(
            (s, sl) for sl in info.slides for s in sl.shapes
            if s.kind not in {"table", "group"}
            and len({r.paragraph for r in s.runs}) >= 2
        )

    def _edit(self, deck, out, shape, slide, after: str):
        apply_changes(deck, approved(deck, Change(
            id="c", op=Op.SET_TEXT, slide=slide.number, target=shape.id,
            before=shape.text, after=after,
        )), out)
        return next(
            s for s in inspect(out).slides[slide.number - 1].shapes if s.id == shape.id
        )

    def test_a_whole_shape_edit_leaves_no_blank_line_behind(self, adversarial_deck, tmp_path):
        shape, slide = self._bulleted(adversarial_deck)
        blank_before = shape.paragraph_count - len({r.paragraph for r in shape.runs})

        after = self._edit(
            adversarial_deck, tmp_path / "out.pptx", shape, slide, "Only this line now"
        )

        blank_after = after.paragraph_count - len({r.paragraph for r in after.runs})
        assert after.text == "Only this line now"
        assert blank_after == blank_before, (
            f"the edit left {blank_after - blank_before} blank line(s) on the slide"
        )

    def test_the_lines_asked_for_are_the_lines_that_arrive(self, adversarial_deck, tmp_path):
        shape, slide = self._bulleted(adversarial_deck)
        after = self._edit(
            adversarial_deck, tmp_path / "out.pptx", shape, slide, "First\nSecond\nThird"
        )
        assert after.text == "First\nSecond\nThird"
        assert len({r.paragraph for r in after.runs}) == 3, (
            "three lines means three paragraphs, not one run holding two newlines"
        )

    def test_the_diff_reports_a_line_that_went(self, adversarial_deck, tmp_path):
        """The second line of defence: the instrument has to be able to see it.

        Fixing the applier is not enough on its own. Every verifier-side lock is
        a question put to the diff, so a property the diff cannot compute is a
        property no lock can protect.
        """
        from slide_wright.diff import diff

        shape, slide = self._bulleted(adversarial_deck)
        out = tmp_path / "out.pptx"
        self._edit(adversarial_deck, out, shape, slide, "Only this line now")

        deltas = [d for d in diff(adversarial_deck, out).deltas if "lines" in d.summary]
        assert deltas, "a shape that lost lines must produce a delta saying so"
        assert not deltas[0].is_content, "a blank line is how it looks, not what it says"

    def test_a_groups_child_keeps_its_one_paragraph(self, adversarial_deck, tmp_path):
        """A group's lines are separate text bodies, and each must keep one.

        Setting a group's text to one line genuinely leaves its other children
        saying nothing; there is no way to express that without deleting shapes,
        which a text edit may not do. So they are emptied, and the diff reports
        the blank lines that result -- which is the truth about the slide.
        """
        from slide_wright.diff import diff

        info = inspect(adversarial_deck)
        group, slide = next(
            (s, sl) for sl in info.slides for s in sl.shapes if s.kind == "group"
        )
        out = tmp_path / "out.pptx"
        after = self._edit(adversarial_deck, out, group, slide, "Only this line now")

        assert after.paragraph_count == group.paragraph_count, (
            "emptying is the only safe answer; a text body must keep a paragraph"
        )
        assert after.text == "Only this line now"
        assert [d.summary for d in diff(adversarial_deck, out).deltas if "blank" in d.summary], (
            "the empty text boxes it now carries must be reported, not hidden"
        )


class TestEveryMultiLineShapeEndsWithTheLinesAskedFor:
    """The property, swept over whatever the decks happen to contain.

    Adding a bulleted body to the corpus caught nothing on its own -- the same
    thing the split-run slide did -- because every test that could see it asks
    about a shape it chose. So this asks the question of every text body in
    reach, five ways: the lines it ends with must be the lines it was asked
    for, it must not have gained a blank one, it must not have invented a link,
    a field or a line break, and no text body may come out with no paragraph at
    all.

    Swept on the elements rather than through `apply_changes`: rewriting the
    package once per shape costs about a second on a 350-part deck, which is
    fifty seconds of suite for arithmetic that lives entirely in the element.
    965 (shape, replacement) pairs over the real corpus run in 0.3s. The package
    round trip is covered end to end above.
    """

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

    #: Five shapes of replacement, because the arithmetic branches on how many
    #: lines arrive against how many were covered: fewer means paragraphs are
    #: removed, more means they are cloned, equal means neither, and empty is
    #: the edge where a text body must still keep one. `None` is "as many lines
    #: as it had", filled per shape.
    REPLACEMENTS = {
        "one line": "Only this",
        "two lines": "Replaced line one\nReplaced line two",
        "eight lines": "\n".join(f"Line {i}" for i in range(1, 9)),
        "empty": "",
        "the same count": None,
    }

    def _shapes(self, path):
        """Every text body on every slide, groups included -- each holds its own."""
        from lxml import etree

        with zipfile.ZipFile(path) as z:
            for name in sorted(z.namelist()):
                if not re.match(r"^ppt/slides/slide\d+\.xml$", name):
                    continue
                root = etree.fromstring(z.read(name))
                for sp in root.iter(f"{self.P}sp"):
                    yield name, sp

    def _lines(self, sp) -> list[str]:
        """Paragraph texts, in order, the way `ShapeInfo.text` assembles them."""
        out = []
        for para in sp.findall(f".//{self.A}p"):
            text = "".join(
                t.text for t in para.findall(f".//{self.A}t") if t.text
            )
            out.append(text)
        return out

    def _count(self, sp, tag: str) -> int:
        return len(sp.findall(f".//{self.A}{tag}"))

    def _sweep(self, path) -> int:
        import copy

        from slide_wright.apply import _set_text

        swept = 0
        for part, original in self._shapes(path):
            lines = self._lines(original)
            written = [line for line in lines if line]
            if len(written) < 2:
                continue
            blank_before = len(lines) - len(written)

            for label, replacement in self.REPLACEMENTS.items():
                sp = copy.deepcopy(original)
                after = replacement
                if after is None:
                    after = "\n".join(f"New {i}" for i in range(len(written)))
                where = f"{path.name} {part} ({label})"
                carried = {
                    tag: self._count(sp, tag) for tag in ("hlinkClick", "fld", "br")
                }

                assert _set_text(sp, "\n".join(written), after), (
                    f"{where}: the joined text could not be found in the shape it "
                    "was read from"
                )

                got = self._lines(sp)
                expected = after.split("\n")
                assert [x for x in got if x] == [x for x in expected if x], (
                    f"{where}: asked for {expected[:3]}, got {got[:3]}"
                )
                assert len(got) == len(expected) + blank_before, (
                    f"{where}: {len(got)} lines where "
                    f"{len(expected) + blank_before} belong"
                )
                for tag, was in carried.items():
                    assert self._count(sp, tag) <= was, (
                        f"{where}: an edit invented a {tag} -- "
                        f"{was} -> {self._count(sp, tag)}"
                    )
                for body in sp.iter(f"{self.A}txBody"):
                    assert body.findall(f"{self.A}p"), (
                        f"{where}: a text body was left with no paragraph at all"
                    )
                swept += 1
        return swept

    def test_the_adversarial_deck(self, adversarial_deck):
        assert self._sweep(adversarial_deck) >= 1, (
            "the corpus has no multi-line text body; the question cannot fail"
        )

    @pytest.mark.fixtures
    def test_the_real_corpus(self):
        fixtures = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "third-party"
        decks = sorted(fixtures.glob("*.pptx"))
        if not decks:
            pytest.skip("third-party corpus not present; run scripts/fetch_fixtures.py")
        swept = sum(self._sweep(deck) for deck in decks)
        assert swept >= 500, f"only {swept} (shape, replacement) pairs swept"
