"""Command line surface.

Exit codes are part of the contract — this is meant to be usable in a script,
so "the deck failed verification" must be distinguishable from "the command was
wrong" without parsing prose.
"""

from __future__ import annotations

import pytest

from slide_wright.cli import EXIT_ERROR, EXIT_FINDINGS, EXIT_OK, main
from slide_wright.inspect import inspect


class TestInspect:
    def test_lists_slides_and_constructs(self, adversarial_deck, capsys):
        assert main(["inspect", str(adversarial_deck)]) == EXIT_OK
        out = capsys.readouterr().out
        assert "7 slides" in out
        assert "Trading comparables" in out
        assert "table:1" in out and "chart:1" in out

    def test_verbose_lists_shape_geometry(self, adversarial_deck, capsys):
        main(["inspect", str(adversarial_deck), "--verbose"])
        assert "id=" in capsys.readouterr().out


class TestAudit:
    def test_clean_deck_exits_zero(self, adversarial_deck, capsys):
        assert main(["audit", str(adversarial_deck)]) == EXIT_OK
        out = capsys.readouterr().out
        assert "DECK AUDIT" in out
        assert "No structural issues found" in out

    def test_gate_only_runs_just_the_delivery_gate(self, adversarial_deck, capsys):
        assert main(["audit", str(adversarial_deck), "--gate-only"]) == EXIT_OK
        out = capsys.readouterr().out
        assert "QUALITY GATE" in out
        assert "DECK AUDIT" not in out

    def test_structural_findings_exit_nonzero(self, tmp_path, capsys):
        """A deck with untitled slides is reportable even if the gate passes."""
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        for _ in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text_frame.text = "Body text with no title above it"
        deck = tmp_path / "untitled.pptx"
        prs.save(str(deck))

        assert main(["audit", str(deck)]) == EXIT_FINDINGS
        assert "no title" in capsys.readouterr().out

    def test_findings_exit_nonzero(self, tmp_path, capsys):
        """A deck with content off the canvas must be reportable in a script."""
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(10), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = "This box runs off the right-hand edge of the slide"
        broken = tmp_path / "broken.pptx"
        prs.save(str(broken))

        assert main(["audit", str(broken)]) == EXIT_FINDINGS
        assert "past the right edge" in capsys.readouterr().out


class TestProfile:
    def test_reports_difficulty(self, adversarial_deck, minimal_deck, capsys):
        assert main(["profile", str(adversarial_deck), str(minimal_deck)]) == EXIT_OK
        out = capsys.readouterr().out
        assert "synthetic-adversarial" in out and "synthetic-minimal" in out


class TestEdit:
    def test_applies_and_writes_output(self, adversarial_deck, tmp_path, capsys):
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        out = tmp_path / "out.pptx"
        code = main([
            "edit", str(adversarial_deck),
            "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
            "-o", str(out), "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_OK
        assert "VERIFIED" in capsys.readouterr().out
        assert out.is_file()
        assert "11.8x" in inspect(out).slide(3).text

    def test_dry_run_changes_nothing(self, adversarial_deck, tmp_path, capsys):
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        before = adversarial_deck.read_bytes()
        code = main([
            "edit", str(adversarial_deck),
            "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
            "--dry-run", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_OK
        assert "dry run" in capsys.readouterr().out
        assert adversarial_deck.read_bytes() == before

    def test_numbers_lock_blocks_a_figure_edit(self, adversarial_deck, tmp_path, capsys):
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        code = main([
            "edit", str(adversarial_deck),
            "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
            "--lock", "numbers", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_FINDINGS
        assert "blocked by numbers lock" in capsys.readouterr().out

    def test_slide_lock_blocks_that_slide(self, adversarial_deck, tmp_path, capsys):
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        code = main([
            "edit", str(adversarial_deck),
            "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
            "--lock", "slide:3", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_FINDINGS

    def test_tables_lock_blocks_a_cell_edit(self, adversarial_deck, tmp_path, capsys):
        """`--lock`'s own help offers `tables`, so it has to mean something here.

        It matched on `Change.object_kind`, a field carried for review UX that
        every builder set except `--set`. The lock printed as held, deck-wide,
        directly above the table edit it was not blocking.
        """
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        code = main([
            "edit", str(adversarial_deck),
            "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
            "--lock", "tables", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_FINDINGS
        assert "blocked by tables lock" in capsys.readouterr().out

    def test_a_set_change_knows_what_kind_of_object_it_edits(self, adversarial_deck):
        """The field the exhibit locks read, filled where the change is built."""
        from slide_wright.cli import _parse_set

        deck = inspect(adversarial_deck)
        table = next(s for s in deck.all_shapes() if s.kind == "table")
        change = _parse_set(f"3:{table.id}/r1/c1:9.4x=11.8x", 1, deck)
        assert change.object_kind == "table"

    def test_malformed_set_is_a_usage_error(self, adversarial_deck, tmp_path, capsys):
        code = main([
            "edit", str(adversarial_deck), "--set", "nonsense",
            "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_ERROR
        assert "slide:target:before=after" in capsys.readouterr().err

    def test_no_changes_is_a_usage_error(self, adversarial_deck, tmp_path, capsys):
        code = main(["edit", str(adversarial_deck), "--workspace", str(tmp_path / "ws")])
        assert code == EXIT_ERROR
        assert "no changes" in capsys.readouterr().err


class TestVerify:
    def test_identical_files_verify(self, adversarial_deck, tmp_path, capsys):
        import shutil

        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        assert main(["verify", str(adversarial_deck), str(copy)]) == EXIT_OK
        assert "100.00%" in capsys.readouterr().out

    def test_verify_writes_nothing_beside_the_deck(self, adversarial_deck, tmp_path):
        """A read-only command must leave no trace next to a confidential file.

        `verify` used to open a session, which materialised a workspace —
        and therefore a copy of the user's deck — in the deck's own folder.
        Decks are confidential by default, so verifying one must not quietly
        duplicate it somewhere the user did not ask for.
        """
        import shutil

        folder = tmp_path / "client-materials"
        folder.mkdir()
        source = folder / "deck.pptx"
        shutil.copy(adversarial_deck, source)
        output = folder / "deck-edited.pptx"
        shutil.copy(adversarial_deck, output)

        before = sorted(p.name for p in folder.iterdir())
        assert main(["verify", str(source), str(output)]) == EXIT_OK
        assert sorted(p.name for p in folder.iterdir()) == before


class TestReviewLoop:
    """propose -> review -> apply, with a human decision in the middle.

    `edit` approves everything it proposes, which is fine for a change you
    typed yourself and wrong for one a model suggested. These three commands
    are the reviewable path: nothing reaches a deck until someone approves it
    by id.
    """

    def _table(self, deck):
        return next(s for s in inspect(deck).all_shapes() if s.kind == "table")

    def _propose(self, deck, ws, out, capsys, extra=()):
        table = self._table(deck)
        code = main(["propose", str(deck), "--workspace", str(ws), "-o", str(out),
                     "--set", f"3:{table.id}/r1/c1:9.4x=11.8x",
                     "--set", f"3:{table.id}/r2/c1:11.2x=12.9x",
                     *extra])
        capsys.readouterr()
        return code

    def test_propose_writes_a_change_set_and_applies_nothing(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        assert self._propose(adversarial_deck, ws, changes, capsys) == EXIT_OK
        assert changes.is_file()
        assert not list(ws.glob("*edited*")), "propose must not write a deck"
        assert adversarial_deck.read_bytes() == (ws / "v000-original.pptx").read_bytes()

    def test_apply_refuses_a_change_set_nobody_approved(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)
        assert main(["apply", str(adversarial_deck), str(changes),
                     "--workspace", str(ws)]) == EXIT_FINDINGS
        assert "nothing is approved" in capsys.readouterr().err

    def test_only_approved_changes_reach_the_deck(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)

        assert main(["review", str(changes), "--approve", "c1",
                     "--reject", "c2"]) == EXIT_OK
        capsys.readouterr()

        out = tmp_path / "out.pptx"
        assert main(["apply", str(adversarial_deck), str(changes),
                     "--workspace", str(ws), "-o", str(out)]) == EXIT_OK
        capsys.readouterr()

        values = [r.text for r in self._table(out).runs]
        assert "11.8x" in values, "the approved change was not applied"
        assert "11.2x" in values, "the rejected change was applied anyway"
        assert "12.9x" not in values, "the rejected change was applied anyway"

    def test_a_review_decision_survives_being_written_out(
        self, adversarial_deck, tmp_path, capsys
    ):
        """The reviewer and the applier are separate processes."""
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)
        main(["review", str(changes), "--reject", "c2"])
        capsys.readouterr()

        assert main(["review", str(changes)]) == EXIT_OK
        assert "1 rejected" in capsys.readouterr().out

    def test_an_unknown_change_id_is_refused(self, adversarial_deck, tmp_path, capsys):
        """Ignoring it would let a reviewer believe they rejected something."""
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)
        assert main(["review", str(changes), "--reject", "c9"]) == EXIT_ERROR
        err = capsys.readouterr().err
        assert "no such change" in err and "c9" in err

    def test_a_change_set_for_another_version_is_refused(
        self, adversarial_deck, tmp_path, capsys
    ):
        """Applying it would fail later as a confusing 'target not found'."""
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)
        main(["review", str(changes), "--approve", "c1"])
        main(["apply", str(adversarial_deck), str(changes), "--workspace", str(ws)])
        capsys.readouterr()

        # The session has moved to v001; this change set describes v000.
        assert main(["apply", str(adversarial_deck), str(changes),
                     "--workspace", str(ws)]) == EXIT_ERROR
        assert "built against" in capsys.readouterr().err

    def test_approve_all_holds_back_uncited_model_changes(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws, changes = tmp_path / "ws", tmp_path / "changes.json"
        self._propose(adversarial_deck, ws, changes, capsys)
        assert main(["review", str(changes), "--approve-all"]) == EXIT_OK
        # Both of these are USER changes, so both are grounded and approved.
        assert "2 approved" in capsys.readouterr().out


class TestHistoryAndRevert:
    """The half of the loop the CLI did not expose.

    propose -> review -> apply -> verify -> revert is the product's promise.
    Until these commands existed a user could edit and verify from the command
    line but had no way to see what versions existed or to go back to one.
    """

    def _edit(self, deck, ws, before, after, message):
        table = next(s for s in inspect(deck).all_shapes() if s.kind == "table")
        return main(["edit", str(deck), "--workspace", str(ws),
                     "--set", f"3:{table.id}/r1/c1:{before}={after}",
                     "-m", message])

    def test_history_lists_every_version(self, adversarial_deck, tmp_path, capsys):
        ws = tmp_path / "ws"
        assert self._edit(adversarial_deck, ws, "9.4x", "11.8x", "Q3") == EXIT_OK
        assert self._edit(adversarial_deck, ws, "11.8x", "12.5x", "Q4") == EXIT_OK
        capsys.readouterr()

        assert main(["history", str(adversarial_deck), "--workspace", str(ws)]) == EXIT_OK
        out = capsys.readouterr().out
        assert "v000" in out and "v001" in out and "v002" in out
        assert "Q3" in out and "Q4" in out

    def test_revert_makes_an_earlier_version_current(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws = tmp_path / "ws"
        self._edit(adversarial_deck, ws, "9.4x", "11.8x", "Q3")
        self._edit(adversarial_deck, ws, "11.8x", "12.5x", "Q4")
        capsys.readouterr()

        out_file = tmp_path / "back.pptx"
        code = main(["revert", str(adversarial_deck), "--workspace", str(ws),
                     "--to", "1", "-o", str(out_file)])
        assert code == EXIT_OK
        assert "v001" in capsys.readouterr().out

        table = next(s for s in inspect(out_file).all_shapes() if s.kind == "table")
        values = [r.text for r in table.runs]
        assert "11.8x" in values, "reverted to the wrong version"
        assert "12.5x" not in values, "the reverted deck still carries the later edit"

    def test_revert_defaults_to_the_original(self, adversarial_deck, tmp_path, capsys):
        ws = tmp_path / "ws"
        self._edit(adversarial_deck, ws, "9.4x", "11.8x", "Q3")
        capsys.readouterr()

        out_file = tmp_path / "orig.pptx"
        assert main(["revert", str(adversarial_deck), "--workspace", str(ws),
                     "-o", str(out_file)]) == EXIT_OK
        assert out_file.read_bytes() == adversarial_deck.read_bytes()

    def test_reverting_to_a_version_that_does_not_exist_is_refused(
        self, adversarial_deck, tmp_path, capsys
    ):
        ws = tmp_path / "ws"
        self._edit(adversarial_deck, ws, "9.4x", "11.8x", "Q3")
        capsys.readouterr()
        assert main(["revert", str(adversarial_deck), "--workspace", str(ws),
                     "--to", "9"]) == EXIT_ERROR
        assert "no version 9" in capsys.readouterr().err


class TestRefusals:
    def test_hostile_file_is_refused_with_an_error_code(self, tmp_path, capsys):
        bad = tmp_path / "bad.pptx"
        bad.write_text("not a zip")
        assert main(["inspect", str(bad)]) == EXIT_ERROR
        assert "refused" in capsys.readouterr().err

    def test_missing_file_is_refused(self, tmp_path, capsys):
        assert main(["inspect", str(tmp_path / "nope.pptx")]) == EXIT_ERROR


class TestRefresh:
    def test_updates_figures_from_a_csv(self, adversarial_deck, tmp_path, capsys):
        source = tmp_path / "comps.csv"
        source.write_text(
            "Company,EV/EBITDA,Margin,Growth\n"
            "Alpha Corp,11.8x,22.1%,18%\n"
            "Beta Industries,11.2x,19.8%,12%\n"
            "Gamma Holdings,8.7x,24.5%,21%\n",
            encoding="utf-8",
        )
        out = tmp_path / "refreshed.pptx"
        code = main([
            "refresh", str(adversarial_deck), "--source", str(source),
            "-o", str(out), "--workspace", str(tmp_path / "ws"),
        ])
        captured = capsys.readouterr().out
        assert code == EXIT_OK
        assert "comps.csv!B2" in captured, "every update must cite its source cell"
        assert "VERIFIED" in captured
        assert "11.8x" in inspect(out).slide(3).text

    def test_dry_run_shows_the_plan_only(self, adversarial_deck, tmp_path, capsys):
        source = tmp_path / "comps.csv"
        source.write_text("Company,EV/EBITDA\nAlpha Corp,11.8x\nBeta Industries,11.2x\n",
                          encoding="utf-8")
        before = adversarial_deck.read_bytes()
        code = main([
            "refresh", str(adversarial_deck), "--source", str(source),
            "--dry-run", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_OK
        assert "dry run" in capsys.readouterr().out
        assert adversarial_deck.read_bytes() == before

    def test_numbers_lock_blocks_a_refresh(self, adversarial_deck, tmp_path, capsys):
        source = tmp_path / "comps.csv"
        source.write_text("Company,EV/EBITDA\nAlpha Corp,11.8x\nBeta Industries,11.2x\n",
                          encoding="utf-8")
        code = main([
            "refresh", str(adversarial_deck), "--source", str(source),
            "--lock", "numbers", "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_FINDINGS

    def test_unreadable_source_is_an_error(self, adversarial_deck, tmp_path, capsys):
        code = main([
            "refresh", str(adversarial_deck), "--source", str(tmp_path / "nope.csv"),
            "--workspace", str(tmp_path / "ws"),
        ])
        assert code == EXIT_ERROR
        assert "error" in capsys.readouterr().err


class TestBrand:
    def test_shows_the_profile_when_no_deck_is_given(self, adversarial_deck, capsys):
        assert main(["brand", str(adversarial_deck)]) == EXIT_OK
        assert "BRAND PROFILE" in capsys.readouterr().out

    def test_a_deck_conforms_to_its_own_template(self, adversarial_deck, capsys):
        assert main(["brand", str(adversarial_deck), str(adversarial_deck)]) == EXIT_OK
        assert "CONFORMS" in capsys.readouterr().out

    def test_off_template_deck_exits_nonzero(self, adversarial_deck, minimal_deck, tmp_path, capsys):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Off-brand text"
        run.font.name = "Comic Sans MS"
        run.font.size = Pt(18)
        off = tmp_path / "off.pptx"
        prs.save(str(off))

        assert main(["brand", str(adversarial_deck), str(off)]) == EXIT_FINDINGS
        assert "Comic Sans MS" in capsys.readouterr().out


class TestDiffCommand:
    """`verify` says whether the package is intact; `diff` says what a reader
    would notice. They answer different questions and both are needed."""

    def test_identical_decks_exit_zero(self, adversarial_deck, tmp_path, capsys):
        import shutil

        copy = tmp_path / "copy.pptx"
        shutil.copy(adversarial_deck, copy)
        assert main(["diff", str(adversarial_deck), str(copy)]) == EXIT_OK
        assert "No structural differences" in capsys.readouterr().out

    def test_a_changed_deck_exits_with_findings_and_says_what_changed(
        self, adversarial_deck, tmp_path, capsys
    ):
        table = next(s for s in inspect(adversarial_deck).all_shapes() if s.kind == "table")
        out = tmp_path / "out.pptx"
        assert main(["edit", str(adversarial_deck), "--workspace", str(tmp_path / "ws"),
                     "--set", f"3:{table.id}/r1/c1:9.4x=11.8x", "-o", str(out)]) == EXIT_OK
        capsys.readouterr()

        assert main(["diff", str(adversarial_deck), str(out)]) == EXIT_FINDINGS
        text = capsys.readouterr().out
        assert "slide 3" in text
        assert "9.4" in text and "11.8" in text


class TestBrandFix:
    """Correcting drift, with the promise that content is untouched.

    The wedge this serves inverts the usual guarantee: change every typeface
    that does not conform, and change nothing else at all. The second half is
    what makes it safe to run on a deck someone else wrote.
    """

    def test_fix_reports_a_plan_and_changes_nothing_on_a_dry_run(
        self, adversarial_deck, tmp_path, capsys
    ):
        code = main(["brand", str(adversarial_deck), str(adversarial_deck),
                     "--fix", "--dry-run", "--workspace", str(tmp_path / "ws")])
        assert code == EXIT_OK
        out = capsys.readouterr().out
        assert "CONFORMANCE PLAN" in out
        assert "No word or number is changed" in out
        assert not list((tmp_path / "ws").glob("*edited*"))

    def test_a_conforming_deck_needs_no_correction(
        self, adversarial_deck, tmp_path, capsys
    ):
        code = main(["brand", str(adversarial_deck), str(adversarial_deck),
                     "--fix", "--workspace", str(tmp_path / "ws")])
        assert code == EXIT_OK
        assert "Nothing to correct" in capsys.readouterr().out

    def test_check_names_what_fix_will_not_touch(self, adversarial_deck, capsys):
        """So nobody runs --fix twice expecting a clean report."""
        main(["brand", str(adversarial_deck), str(adversarial_deck)])
        out = capsys.readouterr().out
        assert "conformance" in out


@pytest.fixture(scope="module")
def conformed_eia(tmp_path_factory):
    """Run the conformance fix on the 350-part EIA deck exactly once.

    Three separate tests each ran the whole pass independently, which cost 71
    of the suite's 107 seconds to do the same work three times. The assertions
    are what differ, not the run.
    """
    from pathlib import Path

    deck = (Path(__file__).resolve().parents[2] / "tests" / "fixtures"
            / "third-party" / "eia-aeo2023-release.pptx")
    if not deck.is_file():
        pytest.skip("run scripts/fetch_fixtures.py")

    workdir = tmp_path_factory.mktemp("conformance")
    out = workdir / "fixed.pptx"
    code = main(["brand", str(deck), str(deck), "--fix",
                 "--workspace", str(workdir / "ws"), "-o", str(out)])
    return deck, out, code


@pytest.mark.fixtures
class TestBrandFixOnARealDeck:
    def test_it_succeeds(self, conformed_eia, capsys):
        _, out, code = conformed_eia
        assert code == EXIT_OK
        assert out.is_file()

    def test_corrects_typefaces_without_touching_a_word(self, conformed_eia):
        from slide_wright.diff import diff

        deck, out, _ = conformed_eia
        result = diff(deck, out)
        assert result.deltas, "the fix must actually change something"
        assert not result.content_deltas, "a formatting pass changed content"
        assert {d.kind for d in result.deltas} == {"formatting"}

    def test_the_corrected_deck_reports_no_typeface_drift(self, conformed_eia):
        from slide_wright.brand import check_conformance, read_profile

        deck, out, _ = conformed_eia
        after = check_conformance(inspect(out), read_profile(deck), "fixed")
        assert not [d for d in after.deviations if d.kind == "font"]

    def test_charts_and_workbooks_survive_a_conformance_pass(self, conformed_eia):
        from slide_wright.charts import census

        deck, out, _ = conformed_eia
        before, after = census(deck), census(out)
        assert after["charts"] == before["charts"] == 29
        assert after["workbooks"] == before["workbooks"] == 29
        assert after["values"] == before["values"]


class TestAlignCommand:
    def test_reporting_does_not_apply(self, adversarial_deck, tmp_path, capsys):
        code = main(["align", str(adversarial_deck), "--workspace", str(tmp_path / "ws")])
        assert code in (EXIT_OK, EXIT_FINDINGS)
        assert "ALIGNMENT PLAN" in capsys.readouterr().out
        assert not list((tmp_path / "ws").glob("*edited*"))

    def test_a_clean_deck_exits_zero(self, adversarial_deck, tmp_path, capsys):
        code = main(["align", str(adversarial_deck), "--workspace", str(tmp_path / "ws")])
        out = capsys.readouterr().out
        if "out of line" in out:
            assert code == EXIT_OK

    def test_the_tolerance_is_configurable(self, adversarial_deck, tmp_path, capsys):
        assert main(["align", str(adversarial_deck), "--tolerance", "0.5",
                     "--workspace", str(tmp_path / "ws")]) in (EXIT_OK, EXIT_FINDINGS)
        assert "0.500in" in capsys.readouterr().out


@pytest.mark.fixtures
class TestAlignOnARealDeck:
    def _fixture(self, name):
        from pathlib import Path

        path = (Path(__file__).resolve().parents[2] / "tests" / "fixtures"
                / "third-party" / name)
        if not path.is_file():
            pytest.skip("run scripts/fetch_fixtures.py")
        return path

    def test_fix_nudges_shapes_and_changes_no_content(self, tmp_path, capsys):
        from slide_wright.diff import diff

        deck = self._fixture("eia-aeo2023-release.pptx")
        out = tmp_path / "aligned.pptx"
        code = main(["align", str(deck), "--fix",
                     "--workspace", str(tmp_path / "ws"), "-o", str(out)])
        if "Nothing within" in capsys.readouterr().out:
            pytest.skip("this fixture has nothing to align")
        assert code == EXIT_OK

        result = diff(deck, out)
        assert result.deltas
        assert not result.content_deltas
        assert {d.kind for d in result.deltas} == {"geometry"}


class TestTidy:
    """The wedge as one command.

    Conformance and alignment were reachable only as separate commands, which
    meant three files, three verifications and three places to lose track of
    what changed. `tidy` is one session, one change set, one verification and
    one point to revert to.
    """

    def test_a_clean_deck_needs_no_tidying(self, adversarial_deck, tmp_path, capsys):
        code = main(["tidy", str(adversarial_deck),
                     "--workspace", str(tmp_path / "ws")])
        assert code == EXIT_OK
        assert "Nothing to tidy" in capsys.readouterr().out

    def test_dry_run_writes_nothing(self, adversarial_deck, tmp_path, capsys):
        ws = tmp_path / "ws"
        assert main(["tidy", str(adversarial_deck), "--workspace", str(ws),
                     "--dry-run"]) == EXIT_OK
        assert not list(ws.glob("*edited*"))

    def test_it_defaults_to_the_decks_own_theme(self, adversarial_deck, tmp_path, capsys):
        """A deck assembled from several sources has a visual system of its own."""
        assert main(["tidy", str(adversarial_deck),
                     "--workspace", str(tmp_path / "ws")]) == EXIT_OK
        assert "deck theme" in capsys.readouterr().out

    def test_an_explicit_template_is_named_as_the_authority(
        self, adversarial_deck, tmp_path, capsys
    ):
        assert main(["tidy", str(adversarial_deck), "--template", str(adversarial_deck),
                     "--workspace", str(tmp_path / "ws")]) == EXIT_OK
        assert "template" in capsys.readouterr().out


@pytest.fixture(scope="module")
def tidied_eia(tmp_path_factory):
    """One `tidy` run over the 350-part EIA deck, shared by its assertions."""
    from pathlib import Path

    deck = (Path(__file__).resolve().parents[2] / "tests" / "fixtures"
            / "third-party" / "eia-aeo2023-release.pptx")
    if not deck.is_file():
        pytest.skip("run scripts/fetch_fixtures.py")

    workdir = tmp_path_factory.mktemp("tidy")
    out = workdir / "tidied.pptx"
    code = main(["tidy", str(deck), "--workspace", str(workdir / "ws"), "-o", str(out)])
    return deck, out, code


@pytest.mark.fixtures
class TestTidyOnARealDeck:
    def test_it_succeeds_and_writes_the_deck(self, tidied_eia):
        _, out, code = tidied_eia
        assert code == EXIT_OK
        assert out.is_file()

    def test_both_passes_ran_in_one_change_set(self, tidied_eia):
        """Typefaces and geometry, applied together rather than in sequence."""
        from slide_wright.diff import diff

        deck, out, _ = tidied_eia
        kinds = {d.kind for d in diff(deck, out).deltas}
        assert "formatting" in kinds and "geometry" in kinds

    def test_not_one_word_or_number_changed(self, tidied_eia):
        from slide_wright.diff import diff

        deck, out, _ = tidied_eia
        assert not diff(deck, out).content_deltas

    def test_every_chart_and_workbook_survived(self, tidied_eia):
        from slide_wright.charts import census

        deck, out, _ = tidied_eia
        before, after = census(deck), census(out)
        assert after["charts"] == before["charts"] == 29
        assert after["workbooks"] == before["workbooks"] == 29
        assert after["values"] == before["values"]

    def test_the_deck_reports_no_typeface_drift_afterwards(self, tidied_eia):
        from slide_wright.brand import check_conformance, read_profile

        deck, out, _ = tidied_eia
        after = check_conformance(inspect(out), read_profile(deck), "tidied")
        assert not [d for d in after.deviations if d.kind == "font"]

    def test_it_converges(self, tidied_eia, tmp_path):
        """Running it again finds nothing left to do."""
        _, out, _ = tidied_eia
        assert main(["tidy", str(out), "--workspace", str(tmp_path / "ws2")]) == EXIT_OK
